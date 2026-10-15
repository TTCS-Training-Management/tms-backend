import base64
import hashlib
import hmac
import json
import os
import unittest
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

_database_path = Path(__file__).resolve().parent / ".sprint1-tests.sqlite"
_database_path.unlink(missing_ok=True)
os.environ["DATABASE_URL"] = f"sqlite:///{_database_path.as_posix()}"

import main


class Sprint1ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(main.app)
        cls.client.__enter__()
        response = cls.client.post(
            "/auth/login",
            json={"email": "admin@tms.local", "password": "Admin@123"},
        )
        cls.admin_headers = {"Authorization": f"Bearer {response.json()['access_token']}"}

    @classmethod
    def tearDownClass(cls):
        cls.client.__exit__(None, None, None)
        main.engine.dispose()
        _database_path.unlink(missing_ok=True)

    def test_permission_matrix_defaults_to_deny(self):
        roles = self.client.get("/auth/roles", headers=self.admin_headers)
        self.assertEqual(
            {role["code"] for role in roles.json()},
            set(main.ROLE_CODES),
        )
        self.assertTrue(main.permission_allows(["admin"], "billing", "full"))
        self.assertTrue(main.permission_allows(["training_manager"], "users", "read"))
        self.assertFalse(main.permission_allows(["training_manager"], "users", "write"))
        self.assertFalse(main.permission_allows(["instructor"], "billing", "read"))
        self.assertTrue(main.permission_allows(["instructor"], "attendance", "write"))
        self.assertFalse(main.permission_allows(["unknown"], "users", "read"))

    def test_relative_sqlite_url_resolves_from_backend_directory(self):
        resolved = main.make_url(main.resolve_database_url("sqlite:///./tms.db"))
        expected = (Path(main.__file__).resolve().parent / "tms.db").resolve()
        self.assertEqual(Path(resolved.database), expected)

    def test_create_search_filter_and_paginate_accounts(self):
        activation_messages = []
        with patch.object(
            main,
            "send_email",
            side_effect=lambda message: activation_messages.append(message) or True,
        ):
            response = self.client.post(
                "/users",
                headers=self.admin_headers,
                json={
                    "full_name": "Nguyễn An",
                    "email": "nguyen.an@example.test",
                    "phone": "0901234567",
                    "roles": ["instructor", "ta"],
                    "assigned_classes": ["TMS-K14", "TMS-K15"],
                },
            )
        self.assertEqual(response.status_code, 201)
        account = response.json()
        self.assertEqual(account["roles"], ["instructor", "ta"])
        self.assertTrue(account["must_change_password"])
        self.assertTrue(account["handover_required"])
        self.assertEqual(activation_messages[0]["To"], "nguyen.an@example.test")
        self.assertEqual(
            activation_messages[0]["Subject"],
            "Thông tin tài khoản của bạn | TMS",
        )
        self.assertIn("Hệ thống Quản lý Đào tạo TMS", str(activation_messages[0]["From"]))
        activation_text = activation_messages[0].get_body(preferencelist=("plain",))
        self.assertIsNotNone(activation_text)
        self.assertIn("mật khẩu tạm thời", activation_text.get_content())
        activation_html = activation_messages[0].get_body(preferencelist=("html",))
        self.assertIsNotNone(activation_html)
        self.assertIn("Đăng nhập hệ thống", activation_html.get_content())

        response = self.client.get(
            "/users?q=0901234567&role=ta&page=1&page_size=20",
            headers=self.admin_headers,
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["total"], 1)
        self.assertEqual(response.json()["items"][0]["email"], "nguyen.an@example.test")

    def test_permission_is_checked_on_server(self):
        with main.SessionLocal() as db:
            manager = main.User(
                full_name="Training Manager",
                email="manager@example.test",
                password_hash=main.hash_password("Manager123"),
                role="training_manager",
            )
            instructor = main.User(
                full_name="Instructor",
                email="instructor@example.test",
                password_hash=main.hash_password("Instructor123"),
                role="instructor",
            )
            db.add_all([manager, instructor])
            db.flush()
            db.add_all(
                [
                    main.UserRoleAssignment(user_id=manager.id, role="training_manager"),
                    main.UserRoleAssignment(user_id=instructor.id, role="instructor"),
                ]
            )
            db.commit()
            manager_id = manager.id
            manager_token = main.encode_token(manager)
            instructor_token = main.encode_token(instructor)

        manager_headers = {"Authorization": f"Bearer {manager_token}"}
        instructor_headers = {"Authorization": f"Bearer {instructor_token}"}
        manager_profile = self.client.get("/auth/me", headers=manager_headers).json()
        instructor_profile = self.client.get("/auth/me", headers=instructor_headers).json()
        self.assertEqual(manager_profile["permissions"]["users"], "R")
        self.assertEqual(manager_profile["permissions"]["billing"], "R")
        self.assertEqual(instructor_profile["permissions"]["attendance"], "W")
        self.assertEqual(instructor_profile["permissions"]["billing"], "-")
        self.assertEqual(self.client.get("/users", headers=manager_headers).status_code, 200)
        self.assertEqual(self.client.get("/users", headers=instructor_headers).status_code, 403)
        self.assertEqual(
            self.client.post(
                "/users",
                headers=manager_headers,
                json={
                    "full_name": "No Admin",
                    "email": "no-admin@example.test",
                    "roles": ["student"],
                },
            ).status_code,
            403,
        )
        updated_roles = self.client.patch(
            f"/users/{manager_id}/role",
            headers=self.admin_headers,
            json={"roles": ["training_manager", "instructor"]},
        )
        self.assertEqual(updated_roles.status_code, 200)
        self.assertEqual(
            self.client.get("/users", headers=manager_headers).status_code, 200
        )
        removed_manager_role = self.client.patch(
            f"/users/{manager_id}/role",
            headers=self.admin_headers,
            json={"roles": ["instructor"]},
        )
        self.assertEqual(removed_manager_role.status_code, 200)
        self.assertEqual(
            self.client.get("/users", headers=manager_headers).status_code, 403
        )

    def test_account_is_not_created_when_activation_email_cannot_be_sent(self):
        with patch.object(main, "send_email", return_value=False):
            response = self.client.post(
                "/users",
                headers=self.admin_headers,
                json={
                    "full_name": "No Email Account",
                    "email": "no-email@example.test",
                    "roles": ["student"],
                },
            )
        self.assertEqual(response.status_code, 503)
        listing = self.client.get(
            "/users?q=no-email@example.test",
            headers=self.admin_headers,
        )
        self.assertEqual(listing.json()["total"], 0)

    def test_lock_requires_reason_and_invalidates_existing_sessions(self):
        with main.SessionLocal() as db:
            staff = main.User(
                full_name="Teacher To Lock",
                email="teacher@example.test",
                password_hash=main.hash_password("Teacher123"),
                role="instructor",
            )
            db.add(staff)
            db.flush()
            db.add(main.UserRoleAssignment(user_id=staff.id, role="instructor"))
            db.add(
                main.UserClassAssignment(user_id=staff.id, class_name="TMS-K16")
            )
            db.commit()
            staff_id = staff.id
            staff_token = main.encode_token(staff)

        response = self.client.patch(
            f"/users/{staff_id}/status",
            headers=self.admin_headers,
            json={"active": False, "reason": ""},
        )
        self.assertEqual(response.status_code, 422)

        response = self.client.patch(
            f"/users/{staff_id}/status",
            headers=self.admin_headers,
            json={"active": False, "reason": "Nghỉ việc"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["handover_required"])
        self.assertEqual(response.json()["assigned_classes"], ["TMS-K16"])
        self.assertEqual(
            self.client.get(
                "/auth/me", headers={"Authorization": f"Bearer {staff_token}"}
            ).status_code,
            401,
        )

    def test_locked_account_cannot_login_and_can_login_after_unlock(self):
        with main.SessionLocal() as db:
            staff = main.User(
                full_name="Lock and Unlock User",
                email="lock-unlock@example.test",
                password_hash=main.hash_password("LockUnlock123"),
                role="instructor",
            )
            db.add(staff)
            db.flush()
            db.add(main.UserRoleAssignment(user_id=staff.id, role="instructor"))
            db.commit()
            staff_id = staff.id

        locked = self.client.patch(
            f"/users/{staff_id}/status",
            headers=self.admin_headers,
            json={"active": False, "reason": "Nghỉ việc"},
        )
        self.assertEqual(locked.status_code, 200)
        rejected_login = self.client.post(
            "/auth/login",
            json={"email": "lock-unlock@example.test", "password": "LockUnlock123"},
        )
        self.assertEqual(rejected_login.status_code, 403)
        self.assertEqual(
            rejected_login.json()["detail"],
            "Tài khoản của bạn đã bị khóa, xin liên hệ với quản trị viên.",
        )

        unlocked = self.client.patch(
            f"/users/{staff_id}/status",
            headers=self.admin_headers,
            json={"active": True},
        )
        self.assertEqual(unlocked.status_code, 200)
        self.assertTrue(unlocked.json()["is_active"])
        self.assertIsNone(unlocked.json()["lock_reason"])
        accepted_login = self.client.post(
            "/auth/login",
            json={"email": "lock-unlock@example.test", "password": "LockUnlock123"},
        )
        self.assertEqual(accepted_login.status_code, 200)

    def test_logout_revokes_session_on_server(self):
        with main.SessionLocal() as db:
            user = main.User(
                full_name="Logout User",
                email="logout@example.test",
                password_hash=main.hash_password("LogoutPass123"),
                role="student",
            )
            db.add(user)
            db.flush()
            db.add(main.UserRoleAssignment(user_id=user.id, role="student"))
            db.commit()
            token = main.encode_token(user)
        headers = {"Authorization": f"Bearer {token}"}
        self.assertEqual(self.client.post("/auth/logout", headers=headers).status_code, 200)
        self.assertEqual(self.client.get("/auth/me", headers=headers).status_code, 401)

    def test_refresh_renews_session_and_logout_invalidates_every_token(self):
        with main.SessionLocal() as db:
            user = main.User(
                full_name="Refresh User",
                email="refresh@example.test",
                password_hash=main.hash_password("RefreshPass123"),
                role="student",
            )
            db.add(user)
            db.flush()
            db.add(main.UserRoleAssignment(user_id=user.id, role="student"))
            db.commit()
        login = self.client.post(
            "/auth/login",
            json={"email": "refresh@example.test", "password": "RefreshPass123"},
        )
        old_token = login.json()["access_token"]
        old_claims = main.decode_token(old_token)
        refreshed = self.client.post(
            "/auth/refresh",
            headers={"Authorization": f"Bearer {old_token}"},
        )
        self.assertEqual(refreshed.status_code, 200)
        new_token = refreshed.json()["access_token"]
        self.assertNotEqual(old_token, new_token)
        new_claims = main.decode_token(new_token)
        self.assertEqual(new_claims["started"], old_claims["started"])
        self.assertEqual(new_claims["max_exp"], old_claims["max_exp"])
        self.assertEqual(
            self.client.get(
                "/auth/me", headers={"Authorization": f"Bearer {new_token}"}
            ).status_code,
            200,
        )
        logged_out = self.client.post(
            "/auth/logout",
            headers={"Authorization": f"Bearer {new_token}"},
        )
        self.assertEqual(logged_out.status_code, 200)
        for token in (old_token, new_token):
            self.assertEqual(
                self.client.get(
                    "/auth/me", headers={"Authorization": f"Bearer {token}"}
                ).status_code,
                401,
            )

    def test_expired_session_is_rejected(self):
        payload = {
            "sub": 1,
            "ver": 0,
            "jti": "expired-test-token",
            "exp": 1,
            "started": 1,
            "last": 1,
            "max_exp": 1,
        }
        body = base64.urlsafe_b64encode(
            json.dumps(payload, separators=(",", ":")).encode()
        ).decode()
        signature = hmac.new(
            main.TOKEN_SECRET.encode(), body.encode(), hashlib.sha256
        ).hexdigest()
        response = self.client.get(
            "/auth/me",
            headers={"Authorization": f"Bearer {body}.{signature}"},
        )
        self.assertEqual(response.status_code, 401)

    def test_admin_cannot_remove_own_admin_role(self):
        admin_id = self.client.get("/auth/me", headers=self.admin_headers).json()["id"]
        response = self.client.patch(
            f"/users/{admin_id}/role",
            headers=self.admin_headers,
            json={"roles": ["instructor"]},
        )
        self.assertEqual(response.status_code, 400)

    def test_delete_account_removes_user_assignments_and_reset_tokens(self):
        with main.SessionLocal() as db:
            user = main.User(
                full_name="Delete User",
                email="delete-user@example.test",
                password_hash=main.hash_password("DeletePass123"),
                role="instructor",
            )
            db.add(user)
            db.flush()
            db.add_all(
                [
                    main.UserRoleAssignment(user_id=user.id, role="instructor"),
                    main.UserClassAssignment(user_id=user.id, class_name="TMS-K14"),
                    main.PasswordResetToken(
                        token_hash=hashlib.sha256(b"delete-user-token").hexdigest(),
                        user_id=user.id,
                        expires_at=datetime.now(timezone.utc).replace(tzinfo=None)
                        + timedelta(minutes=30),
                    ),
                ]
            )
            db.commit()
            user_id = user.id
            token = main.encode_token(user)

        response = self.client.delete(
            f"/users/{user_id}",
            headers=self.admin_headers,
        )
        self.assertEqual(response.status_code, 204)
        with main.SessionLocal() as db:
            self.assertIsNone(db.get(main.User, user_id))
            self.assertEqual(
                db.scalar(
                    select(main.UserRoleAssignment).where(
                        main.UserRoleAssignment.user_id == user_id
                    )
                ),
                None,
            )
            self.assertEqual(
                db.scalar(
                    select(main.UserClassAssignment).where(
                        main.UserClassAssignment.user_id == user_id
                    )
                ),
                None,
            )
            self.assertEqual(
                db.scalar(
                    select(main.PasswordResetToken).where(
                        main.PasswordResetToken.user_id == user_id
                    )
                ),
                None,
            )
        self.assertEqual(
            self.client.get(
                "/auth/me",
                headers={"Authorization": "Bearer " + token},
            ).status_code,
            401,
        )

    def test_admin_cannot_delete_self(self):
        admin_id = self.client.get("/auth/me", headers=self.admin_headers).json()["id"]
        self_delete = self.client.delete(f"/users/{admin_id}", headers=self.admin_headers)
        self.assertEqual(self_delete.status_code, 400, self_delete.text)

    def test_password_reset_is_generic_single_use_and_expires(self):
        with main.SessionLocal() as db:
            reset_user = main.User(
                full_name="Reset User",
                email="reset@example.test",
                password_hash=main.hash_password("ResetPass123"),
                role="student",
            )
            db.add(reset_user)
            db.flush()
            db.add(main.UserRoleAssignment(user_id=reset_user.id, role="student"))
            db.commit()
        messages = []
        with (
            patch.object(
                main,
                "send_email",
                side_effect=lambda message: messages.append(message) or True,
            ),
            patch.object(main.secrets, "token_urlsafe", return_value="one-time-reset"),
        ):
            known = self.client.post(
                "/auth/forgot-password",
                json={"email": "reset@example.test"},
            )
            unknown = self.client.post(
                "/auth/forgot-password",
                json={"email": "nobody@example.test"},
            )
        self.assertEqual(known.json(), unknown.json())
        self.assertIn("liên kết sẽ được gửi", known.json()["message"])
        self.assertIn("30 phút", known.json()["message"])
        self.assertIn("dùng một lần", known.json()["message"])
        self.assertIn(
            "?reset_token=one-time-reset&account=reset%40example.test",
            messages[0].get_body(preferencelist=("plain",)).get_content(),
        )
        self.assertEqual(
            messages[0]["Subject"],
            "Yêu cầu đặt lại mật khẩu | TMS",
        )
        self.assertIn("Hệ thống Quản lý Đào tạo TMS", str(messages[0]["From"]))
        reset_html = messages[0].get_body(preferencelist=("html",))
        self.assertIsNotNone(reset_html)
        self.assertIn("Liên kết sẽ hết hạn sau", reset_html.get_content())
        self.assertIn("Đặt lại mật khẩu", reset_html.get_content())
        with main.SessionLocal() as db:
            reset_record = db.scalar(
                select(main.PasswordResetToken).where(
                    main.PasswordResetToken.token_hash
                    == hashlib.sha256(b"one-time-reset").hexdigest()
                )
            )
            self.assertIsNotNone(reset_record)
        valid_link = self.client.get(
            "/auth/reset-password/validate",
            params={"token": "one-time-reset"},
        )
        self.assertEqual(valid_link.status_code, 200)
        self.assertEqual(valid_link.json()["email"], "reset@example.test")

        response = self.client.post(
            "/auth/reset-password",
            json={"token": "one-time-reset", "new_password": "UpdatedPass123"},
        )
        self.assertEqual(response.status_code, 200)
        with main.SessionLocal() as db:
            db.expire_all()
            reset_record = db.scalar(
                select(main.PasswordResetToken).where(
                    main.PasswordResetToken.token_hash
                    == hashlib.sha256(b"one-time-reset").hexdigest()
                )
            )
            self.assertIsNotNone(reset_record.used_at)
        reused = self.client.post(
            "/auth/reset-password",
            json={"token": "one-time-reset", "new_password": "AnotherPass123"},
        )
        self.assertEqual(reused.status_code, 400)
        reused_link = self.client.get(
            "/auth/reset-password/validate",
            params={"token": "one-time-reset"},
        )
        self.assertEqual(reused_link.status_code, 400)

    def test_change_password_validates_current_password_and_revokes_old_sessions(self):
        with main.SessionLocal() as db:
            user = main.User(
                full_name="Password Change User",
                email="password-change@example.test",
                password_hash=main.hash_password("CurrentPass123"),
                role="student",
            )
            db.add(user)
            db.flush()
            db.add(main.UserRoleAssignment(user_id=user.id, role="student"))
            db.commit()
            old_token = main.encode_token(user)

        headers = {"Authorization": f"Bearer {old_token}"}
        invalid_current = self.client.post(
            "/auth/change-password",
            headers=headers,
            json={
                "current_password": "WrongPass123",
                "new_password": "UpdatedPass123",
            },
        )
        self.assertEqual(invalid_current.status_code, 400)
        changed = self.client.post(
            "/auth/change-password",
            headers=headers,
            json={
                "current_password": "CurrentPass123",
                "new_password": "UpdatedPass123",
            },
        )
        self.assertEqual(changed.status_code, 200)
        new_token = changed.json()["access_token"]
        self.assertEqual(
            self.client.get("/auth/me", headers=headers).status_code,
            401,
        )
        self.assertEqual(
            self.client.get(
                "/auth/me", headers={"Authorization": f"Bearer {new_token}"}
            ).status_code,
            200,
        )
        old_password_login = self.client.post(
            "/auth/login",
            json={
                "email": "password-change@example.test",
                "password": "CurrentPass123",
            },
        )
        new_password_login = self.client.post(
            "/auth/login",
            json={
                "email": "password-change@example.test",
                "password": "UpdatedPass123",
            },
        )
        self.assertEqual(old_password_login.status_code, 401)
        self.assertEqual(new_password_login.status_code, 200)

    def test_expired_reset_token_is_rejected(self):
        with main.SessionLocal() as db:
            reset_user = main.User(
                full_name="Expired Reset User",
                email="expired-reset@example.test",
                password_hash=main.hash_password("ExpiredPass123"),
                role="student",
            )
            db.add(reset_user)
            db.flush()
            db.add(main.UserRoleAssignment(user_id=reset_user.id, role="student"))
            db.add(
                main.PasswordResetToken(
                    token_hash=hashlib.sha256(b"expired-reset").hexdigest(),
                    user_id=reset_user.id,
                    expires_at=datetime.now(timezone.utc).replace(tzinfo=None)
                    - timedelta(seconds=1),
                )
            )
            db.commit()
        response = self.client.post(
            "/auth/reset-password",
            json={"token": "expired-reset", "new_password": "NewPassword123"},
        )
        self.assertEqual(response.status_code, 400)
        expired_link = self.client.get(
            "/auth/reset-password/validate",
            params={"token": "expired-reset"},
        )
        self.assertEqual(expired_link.status_code, 400)

    def test_idle_and_absolute_session_expiry(self):
        now = int(datetime.now(timezone.utc).timestamp())
        for name, claims in (
            (
                "idle",
                {
                    "sub": 1,
                    "ver": 0,
                    "jti": "idle-expired",
                    "started": now - 120,
                    "last": now - 901,
                    "max_exp": now + 1000,
                    "exp": now + 1000,
                },
            ),
            (
                "absolute",
                {
                    "sub": 1,
                    "ver": 0,
                    "jti": "absolute-expired",
                    "started": now - 8 * 60 * 60,
                    "last": now,
                    "max_exp": now - 1,
                    "exp": now + 1000,
                },
            ),
        ):
            with self.subTest(name=name):
                body = base64.urlsafe_b64encode(
                    json.dumps(claims, separators=(",", ":")).encode()
                ).decode()
                signature = hmac.new(
                    main.TOKEN_SECRET.encode(), body.encode(), hashlib.sha256
                ).hexdigest()
                response = self.client.get(
                    "/auth/me",
                    headers={"Authorization": f"Bearer {body}.{signature}"},
                )
                self.assertEqual(response.status_code, 401)
        with main.SessionLocal() as db:
            user = db.scalar(select(main.User).where(main.User.email == "admin@tms.local"))
            self.assertIsNotNone(user)
            idle_token = main.encode_token(user)
        token_body = idle_token.split(".", 1)[0]
        claims = json.loads(base64.urlsafe_b64decode(token_body.encode()))
        claims["last"] = now - 901
        body = base64.urlsafe_b64encode(
            json.dumps(claims, separators=(",", ":")).encode()
        ).decode()
        signature = hmac.new(
            main.TOKEN_SECRET.encode(), body.encode(), hashlib.sha256
        ).hexdigest()
        idle_token = f"{body}.{signature}"
        refreshed = self.client.post(
            "/auth/refresh",
            headers={"Authorization": f"Bearer {idle_token}"},
        )
        self.assertEqual(refreshed.status_code, 401)

    def test_five_bad_passwords_temporarily_lock_account(self):
        statuses = []
        for _ in range(5):
            response = self.client.post(
                "/auth/login",
                json={"email": "accountant@tms.local", "password": "WrongPassword1"},
            )
            statuses.append(response.status_code)
        self.assertEqual(statuses, [401, 401, 401, 401, 423])
        self.assertEqual(response.headers["retry-after"], "900")

        blocked_retry = self.client.post(
            "/auth/login",
            json={"email": "accountant@tms.local", "password": "WrongPassword1"},
        )
        self.assertEqual(blocked_retry.status_code, 423)
        self.assertGreater(int(blocked_retry.headers["retry-after"]), 0)
        with main.SessionLocal() as db:
            account = db.scalar(
                select(main.User).where(main.User.email == "accountant@tms.local")
            )
            self.assertEqual(account.failed_login_attempts, 5)

    def test_successful_login_resets_consecutive_failed_attempts(self):
        email = "reset-failed-attempts@example.test"
        password = "CorrectPass123"
        with main.SessionLocal() as db:
            user = main.User(
                full_name="Failed Attempt Reset User",
                email=email,
                password_hash=main.hash_password(password),
                role="student",
            )
            db.add(user)
            db.flush()
            user_id = user.id
            db.add(main.UserRoleAssignment(user_id=user_id, role="student"))
            db.commit()

        for _ in range(4):
            failed = self.client.post(
                "/auth/login",
                json={"email": email, "password": "WrongPassword1"},
            )
            self.assertEqual(failed.status_code, 401)

        succeeded = self.client.post(
            "/auth/login",
            json={"email": email, "password": password},
        )
        self.assertEqual(succeeded.status_code, 200, succeeded.text)
        with main.SessionLocal() as db:
            db.expire_all()
            user = db.get(main.User, user_id)
            self.assertEqual(user.failed_login_attempts, 0)
            self.assertIsNone(user.locked_until)

        failed_again = self.client.post(
            "/auth/login",
            json={"email": email, "password": "WrongPassword1"},
        )
        self.assertEqual(failed_again.status_code, 401)
        with main.SessionLocal() as db:
            db.expire_all()
            user = db.get(main.User, user_id)
            self.assertEqual(user.failed_login_attempts, 1)

    def test_login_does_not_reveal_whether_email_exists(self):
        known = self.client.post(
            "/auth/login",
            json={"email": "admin@tms.local", "password": "WrongPassword1"},
        )
        unknown = self.client.post(
            "/auth/login",
            json={"email": "missing@example.test", "password": "WrongPassword1"},
        )
        self.assertEqual(known.status_code, 401)
        self.assertEqual(known.json()["detail"], unknown.json()["detail"])

    def test_unknown_and_registered_emails_lock_after_five_failed_logins(self):
        registered_email = "registered-lock-check@example.test"
        unknown_email = "unknown-lock-check@example.test"
        with main.SessionLocal() as db:
            user = main.User(
                full_name="Lock Check User",
                email=registered_email,
                password_hash=main.hash_password("CorrectPassword123"),
                role="student",
            )
            db.add(user)
            db.commit()

        statuses_by_email = {}
        for email in (registered_email, unknown_email):
            statuses = []
            for _ in range(5):
                response = self.client.post(
                    "/auth/login",
                    json={"email": email, "password": "WrongPassword123"},
                )
                statuses.append(response.status_code)
            statuses_by_email[email] = statuses
            self.assertEqual(response.headers["retry-after"], "900")

            blocked_retry = self.client.post(
                "/auth/login",
                json={"email": email, "password": "WrongPassword123"},
            )
            self.assertEqual(blocked_retry.status_code, 423)
            self.assertGreater(int(blocked_retry.headers["retry-after"]), 0)

        self.assertEqual(statuses_by_email[registered_email], [401, 401, 401, 401, 423])
        self.assertEqual(statuses_by_email[unknown_email], [401, 401, 401, 401, 423])
        with main.SessionLocal() as db:
            attempt = db.get(main.LoginAttempt, unknown_email)
            self.assertIsNotNone(attempt)
            self.assertEqual(attempt.failed_attempts, 5)
            self.assertIsNotNone(attempt.locked_until)

    def test_profile_update_avatar_upload_and_serving(self):
        profile = self.client.patch(
            "/profile",
            headers=self.admin_headers,
            json={
                "full_name": "Admin Updated",
                "phone": "0901234567",
                "date_of_birth": "1990-01-02",
                "address": "Thai Nguyen",
            },
        )
        self.assertEqual(profile.status_code, 200)
        self.assertEqual(profile.json()["date_of_birth"], "1990-01-02")

        cleared = self.client.patch(
            "/profile",
            headers=self.admin_headers,
            json={"date_of_birth": None},
        )
        self.assertEqual(cleared.status_code, 200)
        self.assertIsNone(cleared.json()["date_of_birth"])

        uploaded_paths: list[Path] = []
        try:
            uploaded = self.client.post(
                "/profile/avatar",
                headers=self.admin_headers,
                files={"file": ("avatar.png", b"first-image", "image/png")},
            )
            first_avatar_url = uploaded.json()["avatar_url"]
            first_avatar_path = main.UPLOADS_DIR / Path(first_avatar_url).name
            uploaded_paths.append(first_avatar_path)
            self.assertEqual(first_avatar_path.read_bytes(), b"first-image")
            self.assertEqual(self.client.get(first_avatar_url).content, b"first-image")

            replaced = self.client.post(
                "/profile/avatar",
                headers=self.admin_headers,
                files={"file": ("avatar.png", b"replacement-image", "image/png")},
            )
            replacement_url = replaced.json()["avatar_url"]
            uploaded_paths.append(main.UPLOADS_DIR / Path(replacement_url).name)
            self.assertEqual(replaced.status_code, 200)
            self.assertNotEqual(replacement_url, first_avatar_url)
            self.assertFalse(first_avatar_path.exists())
            self.assertEqual(
                self.client.get(replacement_url).content,
                b"replacement-image",
            )
            self.assertEqual(
                self.client.get("/profile", headers=self.admin_headers).json()["avatar_url"],
                replacement_url,
            )
        finally:
            for path in uploaded_paths:
                path.unlink(missing_ok=True)
        self.assertEqual(uploaded.status_code, 200)

    def test_user_import_template_and_row_summary(self):
        template = self.client.get(
            "/users/import-excel/template",
            headers=self.admin_headers,
        )
        self.assertEqual(template.status_code, 200)
        workbook = load_workbook(BytesIO(template.content), read_only=True)
        self.assertEqual(
            list(next(workbook.active.iter_rows(values_only=True))),
            ["full_name", "email", "roles", "phone", "assigned_classes", "initial_password"],
        )
        workbook.close()

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["full_name", "email", "roles", "phone", "assigned_classes", "initial_password"])
        sheet.append(["Import User", "imported@example.test", "student", "0901234567", "Class A", "Temporary123!"])
        sheet.append(["", "missing-name@example.test", "student", "", "", ""])
        sheet.append(["Invalid Email", "not-an-email", "student", "", "", "Temporary123!"])
        sheet.append(["Duplicate User", "imported@example.test", "student", "", "", "Temporary123!"])
        sheet.append(["Invalid Role", "invalid-role@example.test", "student,unknown-role", "", "", "Temporary123!"])
        sheet.append(["Weak Password", "weak-password@example.test", "student", "", "", "weak"])
        for index in range(24):
            sheet.append([f"Import User {index + 2}", f"imported-{index + 2}@example.test", "student", "", "", "Temporary123!"])
        content = BytesIO()
        workbook.save(content)
        with patch.object(main, "smtp_configured", return_value=False):
            preview = self.client.post(
                "/users/import-excel/preview",
                headers=self.admin_headers,
                files={
                    "file": (
                        "users.xlsx",
                        content.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.json()["valid"], 25)
        self.assertEqual(preview.json()["skipped"], 5)
        self.assertEqual(len(preview.json()["preview_rows"]), 25)
        self.assertEqual(
            preview.json()["preview_rows"][0],
            {
                "row_number": 2,
                "full_name": "Import User",
                "email": "imported@example.test",
                "roles": ["student"],
                "phone": "0901234567",
                "assigned_classes": ["Class A"],
            },
        )
        self.assertTrue(all("initial_password" not in row for row in preview.json()["preview_rows"]))
        self.assertTrue(any("Dòng 3" in error for error in preview.json()["errors"]))
        self.assertTrue(any("Dòng 4" in error for error in preview.json()["errors"]))
        self.assertTrue(any("Dòng 5" in error for error in preview.json()["errors"]))
        self.assertTrue(any("Dòng 6" in error for error in preview.json()["errors"]))
        self.assertTrue(any("Dòng 7" in error for error in preview.json()["errors"]))
        with patch.object(main, "smtp_configured", return_value=False):
            imported = self.client.post(
                "/users/import-excel",
                headers=self.admin_headers,
                files={
                    "file": (
                        "users.xlsx",
                        content.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
        self.assertEqual(imported.status_code, 200)
        self.assertEqual(imported.json()["imported"], 25)
        self.assertEqual(imported.json()["skipped"], 5)
        self.assertIn("Dòng 3", imported.json()["errors"][0])
        self.assertEqual(
            self.client.get("/users?q=imported@example.test", headers=self.admin_headers).json()["total"],
            1,
        )

    def test_public_lead_filter_and_assignment(self):
        created = self.client.post(
            "/leads/public",
            json={
                "full_name": "Sprint Lead",
                "phone": "0902222333",
                "email": "lead@example.test",
                "source": "website",
                "program_interest": "TMS",
                "website": "",
            },
        )
        self.assertEqual(created.status_code, 200)

        assignees = self.client.get("/leads/assignees", headers=self.admin_headers)
        self.assertEqual(assignees.status_code, 200)
        admin_id = self.client.get("/auth/me", headers=self.admin_headers).json()["id"]
        self.assertTrue(any(item["id"] == admin_id for item in assignees.json()))

        leads = self.client.get(
            "/leads?q=Sprint+Lead&status=new",
            headers=self.admin_headers,
        )
        self.assertEqual(leads.status_code, 200)
        lead_id = leads.json()[0]["id"]
        paginated = self.client.get(
            "/leads?q=Sprint+Lead&status=new&page=1&page_size=1",
            headers=self.admin_headers,
        )
        self.assertEqual(paginated.status_code, 200, paginated.text)
        self.assertEqual(paginated.json()["total"], 1)
        self.assertEqual(paginated.json()["page"], 1)
        self.assertEqual(paginated.json()["page_size"], 1)
        self.assertEqual(paginated.json()["items"][0]["id"], lead_id)
        assigned = self.client.post(
            f"/leads/{lead_id}/assign",
            headers=self.admin_headers,
            json={"assigned_to_user_id": admin_id},
        )
        self.assertEqual(assigned.status_code, 200)
        changed = self.client.patch(
            f"/leads/{lead_id}",
            headers=self.admin_headers,
            json={"status": "contacted"},
        )
        self.assertEqual(changed.status_code, 200)
        filtered = self.client.get(
            f"/leads?status=contacted&assigned_to={admin_id}",
            headers=self.admin_headers,
        )
        self.assertEqual(len(filtered.json()), 1)
        internal = self.client.post(
            "/leads",
            headers=self.admin_headers,
            json={
                "full_name": "Walk-in Lead",
                "phone": "0902222444",
                "source": "facebook",
            },
        )
        self.assertEqual(internal.status_code, 201, internal.text)
        internal_listing = self.client.get(
            "/leads?q=Walk-in+Lead",
            headers=self.admin_headers,
        )
        self.assertEqual(internal_listing.json()[0]["source"], "facebook")

    def test_lead_list_handles_deleted_or_missing_assignees(self):
        with main.SessionLocal() as db:
            deleted_user = main.User(
                full_name="Deleted Lead Assignee",
                email="deleted-lead-assignee@example.test",
                password_hash=main.hash_password("Deleted123"),
                role="guest",
            )
            db.add(deleted_user)
            db.flush()
            db.add(main.UserRoleAssignment(user_id=deleted_user.id, role="guest"))
            db.add_all([
                main.Lead(
                    full_name="Lead Assigned Before Account Deletion",
                    phone="0907777001",
                    source="website",
                    status="new",
                    assigned_to_user_id=deleted_user.id,
                ),
                main.Lead(
                    full_name="Lead With Historical Missing Assignee",
                    phone="0907777002",
                    source="website",
                    status="new",
                    assigned_to_user_id=999999,
                ),
            ])
            db.commit()
            deleted_user_id = deleted_user.id

        deleted = self.client.delete(
            f"/users/{deleted_user_id}",
            headers=self.admin_headers,
        )
        self.assertEqual(deleted.status_code, 204, deleted.text)

        for lead_name in (
            "Lead Assigned Before Account Deletion",
            "Lead With Historical Missing Assignee",
        ):
            response = self.client.get(
                f"/leads?q={lead_name.replace(' ', '+')}",
                headers=self.admin_headers,
            )
            self.assertEqual(response.status_code, 200, response.text)
            lead = next(item for item in response.json() if item["full_name"] == lead_name)
            self.assertIsNone(lead["assigned_to_user_id"])
            self.assertIsNone(lead["assigned_to_name"])

    def test_lead_crud_duplicate_bulk_assignment_history_and_date_filters(self):
        manager_headers = self._headers_for_role("lead-manager@example.test", "training_manager")
        counselor_headers = self._headers_for_role("lead-counselor@example.test", "admissions")
        self.assertEqual(
            self.client.get("/auth/me", headers=counselor_headers).json()["permissions"]["leads"],
            "W",
        )
        self.assertEqual(
            self.client.get("/auth/me", headers=manager_headers).json()["permissions"]["leads"],
            "F",
        )

        first = self.client.post(
            "/leads",
            headers=self.admin_headers,
            json={
                "full_name": "Bulk Lead One",
                "phone": "0903333001",
                "email": "bulk-one@example.test",
                "source": "website",
                "program_interest": "TMS",
            },
        )
        second = self.client.post(
            "/leads",
            headers=self.admin_headers,
            json={
                "full_name": "Bulk Lead Two",
                "phone": "0903333002",
                "source": "facebook",
            },
        )
        self.assertEqual(first.status_code, 201, first.text)
        self.assertEqual(second.status_code, 201, second.text)
        first_id, second_id = first.json()["id"], second.json()["id"]

        duplicate_check = self.client.get(
            "/leads/duplicate-check",
            headers=counselor_headers,
            params={"phone": "+84903333001"},
        )
        self.assertEqual(duplicate_check.status_code, 200, duplicate_check.text)
        self.assertTrue(duplicate_check.json()["duplicate"])
        duplicate_create = self.client.post(
            "/leads",
            headers=counselor_headers,
            json={"full_name": "Different Person", "phone": "0903333001"},
        )
        self.assertEqual(duplicate_create.status_code, 409)

        bulk = self.client.post(
            "/leads/assign-bulk",
            headers=manager_headers,
            json={"lead_ids": [first_id, second_id], "assigned_to_user_id": self._user_id(counselor_headers)},
        )
        self.assertEqual(bulk.status_code, 200, bulk.text)
        self.assertEqual(bulk.json()["assigned_count"], 2)
        counselor_leads = self.client.get("/leads", headers=counselor_headers)
        self.assertEqual({item["id"] for item in counselor_leads.json()}, {first_id, second_id})
        self.assertEqual(self.client.get(f"/leads/{first_id}/interactions", headers=counselor_headers).status_code, 200)

        edited = self.client.patch(
            f"/leads/{first_id}",
            headers=counselor_headers,
            json={"full_name": "Bulk Lead One Updated", "phone": "0903333001", "email": "updated@example.test"},
        )
        self.assertEqual(edited.status_code, 200, edited.text)
        duplicate_edit = self.client.patch(
            f"/leads/{first_id}",
            headers=counselor_headers,
            json={"phone": "0903333002"},
        )
        self.assertEqual(duplicate_edit.status_code, 409)
        note = self.client.post(
            f"/leads/{first_id}/interactions",
            headers=counselor_headers,
            json={"note": "Đã gọi và hẹn tư vấn ngày mai"},
        )
        self.assertEqual(note.status_code, 201, note.text)
        history = self.client.get(f"/leads/{first_id}/interactions", headers=counselor_headers)
        self.assertTrue(any("hẹn tư vấn ngày mai" in item["note"] for item in history.json()))
        self.assertTrue(any(item["action"] == "assignment" for item in history.json()))
        history_timestamp = next(item["created_at"] for item in history.json() if item["note"] == "Đã gọi và hẹn tư vấn ngày mai")
        self.assertTrue(history_timestamp.endswith("Z"))
        self.assertEqual(datetime.fromisoformat(history_timestamp).utcoffset(), timedelta(0))
        self.assertEqual(self.client.delete(f"/leads/{first_id}", headers=counselor_headers).status_code, 403)

        today = datetime.now(timezone.utc).date().isoformat()
        dated = self.client.get(f"/leads?from_date={today}&to_date={today}", headers=manager_headers)
        self.assertGreaterEqual(len(dated.json()), 2)
        invalid_range = self.client.get(
            "/leads?from_date=2026-10-10&to_date=2026-10-01",
            headers=manager_headers,
        )
        self.assertEqual(invalid_range.status_code, 422)
        self.assertEqual(self.client.delete(f"/leads/{first_id}", headers=self.admin_headers).status_code, 403)
        self.assertEqual(self.client.delete(f"/leads/{first_id}", headers=manager_headers).status_code, 200)
        self.assertEqual(self.client.get(f"/leads/{first_id}/interactions", headers=manager_headers).status_code, 404)

    def _headers_for_role(self, email, role):
        with main.SessionLocal() as db:
            user = main.User(
                full_name=email.split("@")[0],
                email=email,
                password_hash=main.hash_password("LeadTest123"),
                role=role,
            )
            db.add(user)
            db.flush()
            db.add(main.UserRoleAssignment(user_id=user.id, role=role))
            db.commit()
            token = main.encode_token(user)
        return {"Authorization": f"Bearer {token}"}

    def _user_id(self, headers):
        return self.client.get("/auth/me", headers=headers).json()["id"]

    def _create_training_program(self, suffix):
        response = self.client.post(
            "/training/programs",
            headers=self.admin_headers,
            json={"code": f"PRG-{suffix}", "name": f"Program {suffix}"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()["id"]

    def _create_training_subject(self, suffix, session_count=0):
        response = self.client.post(
            "/training/subjects",
            headers=self.admin_headers,
            json={
                "code": f"SUB-{suffix}",
                "name": f"Subject {suffix}",
                "session_count": session_count,
                "weight": 1,
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()["id"]

    def test_program_delete_is_blocked_by_running_class_and_cleans_catalog_links(self):
        program_id = self._create_training_program("CLASS-RULE")
        subject_id = self._create_training_subject("CLASS-RULE")
        linked = self.client.post(
            f"/training/programs/{program_id}/subjects",
            headers=self.admin_headers,
            json={"subject_id": subject_id, "sequence_order": 1},
        )
        self.assertEqual(linked.status_code, 200, linked.text)
        class_response = self.client.post(
            f"/training/programs/{program_id}/classes",
            headers=self.admin_headers,
            json={"name": "Cohort 01"},
        )
        self.assertEqual(class_response.status_code, 201, class_response.text)
        class_id = class_response.json()["id"]
        running = self.client.patch(
            f"/training/classes/{class_id}",
            headers=self.admin_headers,
            json={"status": "running"},
        )
        self.assertEqual(running.status_code, 200, running.text)

        blocked = self.client.delete(
            f"/training/programs/{program_id}",
            headers=self.admin_headers,
        )
        self.assertEqual(blocked.status_code, 409)
        inactive = self.client.patch(
            f"/training/programs/{program_id}",
            headers=self.admin_headers,
            json={"status": "inactive"},
        )
        self.assertEqual(inactive.status_code, 200)
        completed = self.client.patch(
            f"/training/classes/{class_id}",
            headers=self.admin_headers,
            json={"status": "completed"},
        )
        self.assertEqual(completed.status_code, 200)
        deleted = self.client.delete(
            f"/training/programs/{program_id}",
            headers=self.admin_headers,
        )
        self.assertEqual(deleted.status_code, 200, deleted.text)
        self.assertEqual(
            self.client.get(
                f"/training/programs/{program_id}/subjects",
                headers=self.admin_headers,
            ).status_code,
            404,
        )

    def test_training_catalog_search_and_pagination_preserve_legacy_lists(self):
        for suffix in ("PAGE-A", "PAGE-B"):
            program = self.client.post(
                "/training/programs",
                headers=self.admin_headers,
                json={"code": f"CAT-{suffix}", "name": f"Catalog Program {suffix}"},
            )
            self.assertEqual(program.status_code, 200, program.text)
            subject = self.client.post(
                "/training/subjects",
                headers=self.admin_headers,
                json={"code": f"CAT-{suffix}", "name": f"Catalog Subject {suffix}"},
            )
            self.assertEqual(subject.status_code, 200, subject.text)

        program_page = self.client.get(
            "/training/programs?q=Catalog+Program&page=1&page_size=1",
            headers=self.admin_headers,
        )
        self.assertEqual(program_page.status_code, 200, program_page.text)
        self.assertEqual(program_page.json()["total"], 2)
        self.assertEqual(len(program_page.json()["items"]), 1)
        self.assertEqual(program_page.json()["page"], 1)

        subject_page = self.client.get(
            "/training/subjects?q=Catalog+Subject&page=2&page_size=1",
            headers=self.admin_headers,
        )
        self.assertEqual(subject_page.status_code, 200, subject_page.text)
        self.assertEqual(subject_page.json()["total"], 2)
        self.assertEqual(len(subject_page.json()["items"]), 1)
        self.assertEqual(subject_page.json()["page"], 2)

        legacy_programs = self.client.get(
            "/training/programs",
            headers=self.admin_headers,
        )
        legacy_subjects = self.client.get(
            "/training/subjects",
            headers=self.admin_headers,
        )
        self.assertIsInstance(legacy_programs.json(), list)
        self.assertIsInstance(legacy_subjects.json(), list)

    def test_curriculum_rejects_duplicates_and_foreign_prerequisites_and_persists_order(self):
        program_id = self._create_training_program("ORDER")
        other_program_id = self._create_training_program("OTHER")
        subject_a = self._create_training_subject("ORDER-A")
        subject_b = self._create_training_subject("ORDER-B")
        subject_c = self._create_training_subject("ORDER-C")
        subject_d = self._create_training_subject("ORDER-D")
        foreign_subject = self._create_training_subject("FOREIGN")
        linked_a = self.client.post(
            f"/training/programs/{program_id}/subjects",
            headers=self.admin_headers,
            json={"subject_id": subject_a, "sequence_order": 1},
        )
        self.assertEqual(linked_a.status_code, 200, linked_a.text)
        linked_b = self.client.post(
            f"/training/programs/{program_id}/subjects",
            headers=self.admin_headers,
            json={
                "subject_id": subject_b,
                "sequence_order": 2,
                "prerequisite_subject_id": subject_a,
            },
        )
        self.assertEqual(linked_b.status_code, 200, linked_b.text)
        linked_c = self.client.post(
            f"/training/programs/{program_id}/subjects",
            headers=self.admin_headers,
            json={"subject_id": subject_c, "sequence_order": 3},
        )
        self.assertEqual(linked_c.status_code, 200, linked_c.text)
        duplicate = self.client.post(
            f"/training/programs/{program_id}/subjects",
            headers=self.admin_headers,
            json={"subject_id": subject_a, "sequence_order": 4},
        )
        self.assertEqual(duplicate.status_code, 409)
        foreign = self.client.post(
            f"/training/programs/{other_program_id}/subjects",
            headers=self.admin_headers,
            json={"subject_id": foreign_subject, "sequence_order": 1},
        )
        self.assertEqual(foreign.status_code, 200)
        invalid_prerequisite = self.client.post(
            f"/training/programs/{program_id}/subjects",
            headers=self.admin_headers,
            json={
                "subject_id": subject_d,
                "sequence_order": 4,
                "prerequisite_subject_id": foreign_subject,
            },
        )
        self.assertEqual(invalid_prerequisite.status_code, 422)
        self_prerequisite = self.client.post(
            f"/training/programs/{program_id}/subjects",
            headers=self.admin_headers,
            json={
                "subject_id": foreign_subject,
                "sequence_order": 4,
                "prerequisite_subject_id": foreign_subject,
            },
        )
        self.assertEqual(self_prerequisite.status_code, 422)

        rows = self.client.get(
            f"/training/programs/{program_id}/subjects",
            headers=self.admin_headers,
        ).json()
        reordered = self.client.put(
            f"/training/programs/{program_id}/subjects/order",
            headers=self.admin_headers,
            json={"association_ids": [rows[2]["id"], rows[0]["id"], rows[1]["id"]]},
        )
        self.assertEqual(reordered.status_code, 200, reordered.text)
        persisted = self.client.get(
            f"/training/programs/{program_id}/subjects",
            headers=self.admin_headers,
        ).json()
        self.assertEqual(
            [row["subject_id"] for row in persisted],
            [subject_c, subject_a, subject_b],
        )
        removed = self.client.delete(
            f"/training/programs/{program_id}/subjects/{rows[0]['id']}",
            headers=self.admin_headers,
        )
        self.assertEqual(removed.status_code, 200, removed.text)
        remaining = self.client.get(
            f"/training/programs/{program_id}/subjects",
            headers=self.admin_headers,
        ).json()
        subject_b_link = next(row for row in remaining if row["subject_id"] == subject_b)
        self.assertIsNone(subject_b_link["prerequisite_subject_id"])

    def test_subject_cannot_be_deleted_while_used_by_a_running_class(self):
        program_id = self._create_training_program("SUBJECT-LOCK")
        subject_id = self._create_training_subject("SUBJECT-LOCK")
        self.assertEqual(
            self.client.post(
                f"/training/programs/{program_id}/subjects",
                headers=self.admin_headers,
                json={"subject_id": subject_id, "sequence_order": 1},
            ).status_code,
            200,
        )
        training_class = self.client.post(
            f"/training/programs/{program_id}/classes",
            headers=self.admin_headers,
            json={"name": "Running cohort"},
        ).json()
        self.client.patch(
            f"/training/classes/{training_class['id']}",
            headers=self.admin_headers,
            json={"status": "running"},
        )
        blocked = self.client.delete(
            f"/training/subjects/{subject_id}",
            headers=self.admin_headers,
        )
        self.assertEqual(blocked.status_code, 409)
        self.client.patch(
            f"/training/classes/{training_class['id']}",
            headers=self.admin_headers,
            json={"status": "completed"},
        )
        deleted = self.client.delete(
            f"/training/subjects/{subject_id}",
            headers=self.admin_headers,
        )
        self.assertEqual(deleted.status_code, 200, deleted.text)

    def test_subject_sessions_are_unique_bounded_and_clone_in_order(self):
        source_id = self._create_training_subject("SESSION-SOURCE", session_count=2)
        target_id = self._create_training_subject("SESSION-TARGET", session_count=3)
        for sequence in (1, 2):
            created = self.client.post(
                f"/training/subjects/{source_id}/sessions",
                headers=self.admin_headers,
                json={
                    "sequence": sequence,
                    "title": f"Source session {sequence}",
                    "objectives": f"Objective {sequence}",
                },
            )
            self.assertEqual(created.status_code, 200, created.text)
        target_first = self.client.post(
            f"/training/subjects/{target_id}/sessions",
            headers=self.admin_headers,
            json={"sequence": 1, "title": "Existing target session"},
        )
        self.assertEqual(target_first.status_code, 200, target_first.text)
        cloned = self.client.post(
            f"/training/subjects/{target_id}/sessions/clone",
            headers=self.admin_headers,
            json={"source_subject_id": source_id},
        )
        self.assertEqual(cloned.status_code, 200, cloned.text)
        self.assertEqual(cloned.json()["copied_count"], 2)
        target_sessions = self.client.get(
            f"/training/subjects/{target_id}/sessions",
            headers=self.admin_headers,
        ).json()
        self.assertEqual([item["sequence"] for item in target_sessions], [1, 2, 3])
        self.assertEqual(target_sessions[1]["title"], "Source session 1")
        over_limit_clone = self.client.post(
            f"/training/subjects/{target_id}/sessions/clone",
            headers=self.admin_headers,
            json={"source_subject_id": source_id},
        )
        self.assertEqual(over_limit_clone.status_code, 409)
        duplicate_sequence = self.client.post(
            f"/training/subjects/{target_id}/sessions",
            headers=self.admin_headers,
            json={"sequence": 3, "title": "Duplicate"},
        )
        self.assertEqual(duplicate_sequence.status_code, 409)
        out_of_range = self.client.post(
            f"/training/subjects/{target_id}/sessions",
            headers=self.admin_headers,
            json={"sequence": 4, "title": "Over limit"},
        )
        self.assertEqual(out_of_range.status_code, 422)

    def test_public_lead_honeypot_and_rate_limit(self):
        with main.SessionLocal() as db:
            db.execute(main.delete(main.PublicLeadAttempt))
            db.commit()
        trapped = self.client.post(
            "/leads/public",
            json={
                "full_name": "Bot",
                "phone": "0907654300",
                "website": "spam.example",
            },
        )
        self.assertEqual(trapped.status_code, 400)
        for index in range(5):
            accepted = self.client.post(
                "/leads/public",
                json={
                    "full_name": f"Rate Test {index}",
                    "phone": f"09076543{index:02d}",
                    "website": "",
                },
            )
            self.assertEqual(accepted.status_code, 200, accepted.text)
            self.assertIn("sớm liên hệ", accepted.json()["message"])
        limited = self.client.post(
            "/leads/public",
            json={
                "full_name": "Rate Test Limit",
                "phone": "0907654399",
                "website": "",
            },
        )
        self.assertEqual(limited.status_code, 429)
        self.assertEqual(limited.headers.get("retry-after"), "600")


if __name__ == "__main__":
    unittest.main()
