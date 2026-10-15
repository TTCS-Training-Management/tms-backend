from __future__ import annotations

import base64
import hashlib
import hmac
import json
import logging
import os
import re
import secrets
import smtplib
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from email.message import EmailMessage
from email.utils import formataddr, parseaddr
from html import escape
from pathlib import Path
from typing import Annotated, Any, Literal
from urllib.parse import quote
from zipfile import BadZipFile

from dotenv import load_dotenv
from fastapi import Depends, FastAPI, File, HTTPException, Query, Request, Response, UploadFile, status
from fastapi.responses import JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from pydantic import BaseModel, ConfigDict, Field, field_validator
from sqlalchemy import Boolean, Date, DateTime, ForeignKey, Integer, String, create_engine, delete, func, inspect, select, text, update
from sqlalchemy.engine import make_url
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column, sessionmaker
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

BACKEND_DIR = Path(__file__).resolve().parent
load_dotenv(BACKEND_DIR / ".env")
logger = logging.getLogger(__name__)
UPLOADS_DIR = BACKEND_DIR / "uploads" / "avatars"
UPLOADS_DIR.mkdir(parents=True, exist_ok=True)

def resolve_database_url(database_url: str) -> str:
    parsed_url = make_url(database_url)
    database_path = parsed_url.database
    if (
        parsed_url.drivername.startswith("sqlite")
        and database_path not in (None, "", ":memory:")
        and not database_path.startswith("file:")
        and not Path(database_path).is_absolute()
    ):
        absolute_path = (BACKEND_DIR / database_path).resolve()
        return parsed_url.set(database=str(absolute_path)).render_as_string(
            hide_password=False
        )
    return database_url


DATABASE_URL = resolve_database_url(os.getenv("DATABASE_URL", "sqlite:///./tms.db"))
TOKEN_SECRET = os.getenv("TOKEN_SECRET", "tms-local-development-secret")
TOKEN_TTL_HOURS = 8
SESSION_IDLE_MINUTES = 15
SMTP_HOST = os.getenv("SMTP_HOST", "")
SMTP_PORT = int(os.getenv("SMTP_PORT", "587"))
SMTP_USER = os.getenv("SMTP_USER", "")
SMTP_PASSWORD = os.getenv("SMTP_PASSWORD", "")
SMTP_FROM = os.getenv("SMTP_FROM", SMTP_USER)
FRONTEND_URL = os.getenv("FRONTEND_URL", "http://localhost:5173")
ALLOWED_ORIGINS = list(
    {
        FRONTEND_URL,
        "http://localhost:5173",
        "http://localhost:5174",
        "http://127.0.0.1:5173",
        "http://127.0.0.1:5174",
    }
)

RoleCode = Literal[
    "guest",
    "student",
    "instructor",
    "ta",
    "training_manager",
    "admissions",
    "accountant",
    "admin",
]
ROLE_CODES: tuple[str, ...] = (
    "guest",
    "student",
    "instructor",
    "ta",
    "training_manager",
    "admissions",
    "accountant",
    "admin",
)
ROLE_LABELS: dict[str, str] = {
    "guest": "Khách truy cập",
    "student": "Học viên",
    "instructor": "Giảng viên",
    "ta": "Trợ giảng",
    "training_manager": "Quản lý đào tạo",
    "admissions": "Tư vấn tuyển sinh",
    "accountant": "Kế toán",
    "admin": "Quản trị hệ thống",
}

# The matrix follows the User Roles sheet. Unlisted module access is denied.
ROLE_PERMISSIONS: dict[str, dict[str, str]] = {
    "guest": {},
    "student": {
        "courses": "R",
        "student_records": "W",
        "classes": "R",
        "attendance": "W",
        "materials": "R",
        "assignments": "W",
        "grades": "R",
        "billing": "R",
        "surveys": "W",
    },
    "ta": {
        "courses": "R",
        "student_records": "R",
        "classes": "R",
        "attendance": "W",
        "materials": "W",
        "assignments": "W",
        "grades": "R",
    },
    "instructor": {
        "courses": "R",
        "student_records": "R",
        "classes": "R",
        "attendance": "W",
        "materials": "W",
        "assignments": "F",
        "grades": "W",
        "surveys": "R",
        "reports": "R",
    },
    "admissions": {
        "courses": "R",
        "leads": "W",
        "student_records": "W",
        "classes": "R",
        "reports": "R",
    },
    "accountant": {
        "leads": "R",
        "student_records": "R",
        "billing": "F",
        "reports": "R",
    },
    "training_manager": {
        "courses": "F",
        "leads": "F",
        "student_records": "F",
        "classes": "F",
        "attendance": "F",
        "materials": "F",
        "assignments": "R",
        "grades": "F",
        "billing": "R",
        "surveys": "F",
        "reports": "F",
        "users": "R",
    },
    "admin": {
        module: "F"
        for module in (
            "courses",
            "leads",
            "student_records",
            "classes",
            "attendance",
            "materials",
            "assignments",
            "grades",
            "billing",
            "surveys",
            "reports",
            "users",
        )
    },
}
PERMISSION_LEVELS = {"-": 0, "R": 1, "W": 2, "F": 3}
MODULE_CODES = tuple(
    dict.fromkeys(module for permissions in ROLE_PERMISSIONS.values() for module in permissions)
)

connect_args = {"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {}
engine = create_engine(DATABASE_URL, connect_args=connect_args)
SessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)


class Base(DeclarativeBase):
    pass


class User(Base):
    __tablename__ = "users"

    id: Mapped[int] = mapped_column(primary_key=True)
    full_name: Mapped[str] = mapped_column(String(120))
    email: Mapped[str] = mapped_column(String(255), unique=True, index=True)
    password_hash: Mapped[str] = mapped_column(String(128))
    role: Mapped[str] = mapped_column(String(30), default="instructor")
    is_active: Mapped[bool] = mapped_column(Boolean, default=True)
    failed_login_attempts: Mapped[int] = mapped_column(Integer, default=0)
    locked_until: Mapped[datetime | None] = mapped_column(DateTime, nullable=True)
    created_at: Mapped[datetime] = mapped_column(
        DateTime, default=lambda: datetime.now(timezone.utc)
    )
    phone: Mapped[str | None] = mapped_column(String(40), nullable=True)
    date_of_birth: Mapped[date | None] = mapped_column(Date, nullable=True)
    address: Mapped[str | None] = mapped_column(String(500), nullable=True)
    avatar_url: Mapped[str | None] = mapped_column(String(500), nullable=True)
    lock_reason: Mapped[str | None] = mapped_column(String(500), nullable=True)
    auth_version: Mapped[int] = mapped_column(Integer, default=0)
    must_change_password: Mapped[bool] = mapped_column(Boolean, default=False)


class LoginAttempt(Base):
    __tablename__ = "login_attempts"

    email: Mapped[str] = mapped_column(String(255), primary_key=True)
    failed_attempts: Mapped[int] = mapped_column(Integer, default=0)
    locked_until: Mapped[datetime | None] = mapped_column(DateTime, nullable=True)


class UserRoleAssignment(Base):
    __tablename__ = "user_roles"

    user_id: Mapped[int] = mapped_column(
        ForeignKey("users.id", ondelete="CASCADE"), primary_key=True
    )
    role: Mapped[str] = mapped_column(String(30), primary_key=True)


class UserClassAssignment(Base):
    __tablename__ = "user_class_assignments"

    id: Mapped[int] = mapped_column(primary_key=True)
    user_id: Mapped[int] = mapped_column(
        ForeignKey("users.id", ondelete="CASCADE"), index=True
    )
    class_name: Mapped[str] = mapped_column(String(120))


class TrainingProgram(Base):
    __tablename__ = "training_programs"

    id: Mapped[int] = mapped_column(primary_key=True)
    code: Mapped[str] = mapped_column(String(30), unique=True, index=True)
    name: Mapped[str] = mapped_column(String(200))
    description: Mapped[str | None] = mapped_column(String(2000), nullable=True)
    total_duration_hours: Mapped[int] = mapped_column(Integer, default=0)
    standard_fee: Mapped[int] = mapped_column(Integer, default=0)
    status: Mapped[str] = mapped_column(String(20), default="active")
    created_at: Mapped[datetime] = mapped_column(DateTime, default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(DateTime, default=lambda: datetime.now(timezone.utc), onupdate=lambda: datetime.now(timezone.utc))


class TrainingClass(Base):
    __tablename__ = "training_classes"

    id: Mapped[int] = mapped_column(primary_key=True)
    program_id: Mapped[int] = mapped_column(ForeignKey("training_programs.id", ondelete="RESTRICT"), index=True)
    name: Mapped[str] = mapped_column(String(200))
    status: Mapped[str] = mapped_column(String(20), default="planned")
    created_at: Mapped[datetime] = mapped_column(DateTime, default=lambda: datetime.now(timezone.utc))


class TrainingSubject(Base):
    __tablename__ = "training_subjects"

    id: Mapped[int] = mapped_column(primary_key=True)
    code: Mapped[str] = mapped_column(String(30), unique=True, index=True)
    name: Mapped[str] = mapped_column(String(200))
    session_count: Mapped[int] = mapped_column(Integer, default=0)
    weight: Mapped[int] = mapped_column(Integer, default=1)
    description: Mapped[str | None] = mapped_column(String(2000), nullable=True)
    learning_outcomes: Mapped[str | None] = mapped_column(String(2000), nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=lambda: datetime.now(timezone.utc))


class ProgramSubject(Base):
    __tablename__ = "program_subjects"

    id: Mapped[int] = mapped_column(primary_key=True)
    program_id: Mapped[int] = mapped_column(ForeignKey("training_programs.id", ondelete="CASCADE"), index=True)
    subject_id: Mapped[int] = mapped_column(ForeignKey("training_subjects.id", ondelete="CASCADE"), index=True)
    sequence_order: Mapped[int] = mapped_column(Integer, default=1)
    prerequisite_subject_id: Mapped[int | None] = mapped_column(ForeignKey("training_subjects.id", ondelete="SET NULL"), nullable=True)
    is_required: Mapped[bool] = mapped_column(Boolean, default=True)


class SubjectSession(Base):
    __tablename__ = "subject_sessions"

    id: Mapped[int] = mapped_column(primary_key=True)
    subject_id: Mapped[int] = mapped_column(ForeignKey("training_subjects.id", ondelete="CASCADE"), index=True)
    sequence: Mapped[int] = mapped_column(Integer, default=1)
    title: Mapped[str] = mapped_column(String(200))
    objectives: Mapped[str | None] = mapped_column(String(2000), nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=lambda: datetime.now(timezone.utc))


class PublicLeadAttempt(Base):
    __tablename__ = "public_lead_attempts"

    id: Mapped[int] = mapped_column(primary_key=True)
    client_hash: Mapped[str] = mapped_column(String(64), index=True)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=lambda: datetime.now(timezone.utc), index=True)


class Lead(Base):
    __tablename__ = "leads"

    id: Mapped[int] = mapped_column(primary_key=True)
    full_name: Mapped[str] = mapped_column(String(120))
    phone: Mapped[str] = mapped_column(String(40))
    email: Mapped[str | None] = mapped_column(String(255), nullable=True)
    source: Mapped[str] = mapped_column(String(40), default="website")
    status: Mapped[str] = mapped_column(String(30), default="new")
    program_interest: Mapped[str | None] = mapped_column(String(200), nullable=True)
    notes: Mapped[str | None] = mapped_column(String(4000), nullable=True)
    loss_reason: Mapped[str | None] = mapped_column(String(1000), nullable=True)
    converted_at: Mapped[datetime | None] = mapped_column(DateTime, nullable=True)
    assigned_to_user_id: Mapped[int | None] = mapped_column(ForeignKey("users.id", ondelete="SET NULL"), nullable=True, index=True)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(DateTime, default=lambda: datetime.now(timezone.utc), onupdate=lambda: datetime.now(timezone.utc))


class LeadInteraction(Base):
    __tablename__ = "lead_interactions"

    id: Mapped[int] = mapped_column(primary_key=True)
    lead_id: Mapped[int] = mapped_column(ForeignKey("leads.id", ondelete="CASCADE"), index=True)
    user_id: Mapped[int | None] = mapped_column(ForeignKey("users.id", ondelete="SET NULL"), nullable=True, index=True)
    action: Mapped[str] = mapped_column(String(40))
    note: Mapped[str] = mapped_column(String(2000))
    interaction_type: Mapped[str] = mapped_column(String(20), default="message")
    follow_up_at: Mapped[datetime | None] = mapped_column(DateTime, nullable=True, index=True)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=lambda: datetime.now(timezone.utc), index=True)


class PasswordResetToken(Base):
    __tablename__ = "password_reset_tokens"

    id: Mapped[int] = mapped_column(primary_key=True)
    token_hash: Mapped[str] = mapped_column(String(64), unique=True, index=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True)
    expires_at: Mapped[datetime] = mapped_column(DateTime)
    used_at: Mapped[datetime | None] = mapped_column(DateTime, nullable=True)


def validate_password_strength(password: str) -> str:
    if len(password) < 8:
        raise ValueError("Mật khẩu phải có ít nhất 8 ký tự")
    if not re.search(r"[A-Za-z]", password) or not re.search(r"\d", password):
        raise ValueError("Mật khẩu phải có ít nhất 1 chữ cái và 1 chữ số")
    return password


def validate_email(email: str) -> str:
    normalized = email.strip().lower()
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", normalized):
        raise ValueError("Email không hợp lệ")
    return normalized


class LoginRequest(BaseModel):
    email: str
    password: str


class PasswordChangeRequest(BaseModel):
    current_password: str
    new_password: str

    @field_validator("new_password")
    @classmethod
    def validate_new_password(cls, value: str) -> str:
        return validate_password_strength(value)


class ForgotPasswordRequest(BaseModel):
    email: str

    @field_validator("email")
    @classmethod
    def validate_address(cls, value: str) -> str:
        return validate_email(value)


class ResetPasswordRequest(BaseModel):
    token: str
    new_password: str

    @field_validator("new_password")
    @classmethod
    def validate_new_password(cls, value: str) -> str:
        return validate_password_strength(value)


class UserCreate(BaseModel):
    full_name: str = Field(min_length=2, max_length=120)
    email: str
    phone: str | None = Field(default=None, max_length=40)
    roles: list[RoleCode] = Field(min_length=1)
    assigned_classes: list[str] = Field(default_factory=list)

    @field_validator("email")
    @classmethod
    def validate_address(cls, value: str) -> str:
        return validate_email(value)

    @field_validator("roles")
    @classmethod
    def validate_unique_roles(cls, value: list[RoleCode]) -> list[RoleCode]:
        if len(value) != len(set(value)):
            raise ValueError("Vai trò không được trùng lặp")
        return value


class UserUpdate(BaseModel):
    full_name: str = Field(min_length=2, max_length=120)
    email: str
    phone: str | None = Field(default=None, max_length=40)
    assigned_classes: list[str] = Field(default_factory=list)

    @field_validator("email")
    @classmethod
    def validate_address(cls, value: str) -> str:
        return validate_email(value)


class UserRoleUpdate(BaseModel):
    roles: list[RoleCode] = Field(min_length=1)

    @field_validator("roles")
    @classmethod
    def validate_unique_roles(cls, value: list[RoleCode]) -> list[RoleCode]:
        if len(value) != len(set(value)):
            raise ValueError("Vai trò không được trùng lặp")
        return value


class UserStatusUpdate(BaseModel):
    active: bool
    reason: str | None = Field(default=None, max_length=500)


class UserResponse(BaseModel):
    model_config = ConfigDict(from_attributes=True)

    id: int
    full_name: str
    email: str
    phone: str | None
    assigned_classes: list[str]
    handover_required: bool
    role: RoleCode
    roles: list[RoleCode]
    permissions: dict[str, str]
    is_active: bool
    lock_reason: str | None
    must_change_password: bool
    date_of_birth: date | None = None
    address: str | None = None
    avatar_url: str | None = None
    created_at: datetime


class UserPage(BaseModel):
    items: list[UserResponse]
    total: int
    page: int
    page_size: int


class LoginResponse(BaseModel):
    access_token: str
    token_type: str = "bearer"
    user: UserResponse


bearer = HTTPBearer(auto_error=False)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


DbSession = Annotated[Session, Depends(get_db)]
CurrentCredentials = Annotated[
    HTTPAuthorizationCredentials | None, Depends(bearer)
]


def assigned_roles(db: Session, user: User) -> list[str]:
    roles = list(
        db.scalars(
            select(UserRoleAssignment.role)
            .where(UserRoleAssignment.user_id == user.id)
            .order_by(UserRoleAssignment.role)
        ).all()
    )
    if roles:
        return roles
    return [user.role] if user.role in ROLE_CODES else []


def user_response(db: Session, user: User) -> UserResponse:
    roles = assigned_roles(db, user)
    primary_role = user.role if user.role in roles else (roles[0] if roles else "guest")
    permissions = {
        module: max(
            (
                ROLE_PERMISSIONS.get(role, {}).get(module, "-")
                for role in roles
            ),
            key=lambda level: PERMISSION_LEVELS[level],
            default="-",
        )
        for module in MODULE_CODES
    }
    assigned_classes = list(
        db.scalars(
            select(UserClassAssignment.class_name)
            .where(UserClassAssignment.user_id == user.id)
            .order_by(UserClassAssignment.class_name)
        ).all()
    )
    return UserResponse(
        id=user.id,
        full_name=user.full_name,
        email=user.email,
        phone=user.phone,
        assigned_classes=assigned_classes,
        handover_required=bool(assigned_classes),
        role=primary_role,
        roles=roles,
        permissions=permissions,
        is_active=user.is_active,
        lock_reason=user.lock_reason,
        must_change_password=user.must_change_password,
        date_of_birth=getattr(user, "date_of_birth", None),
        address=getattr(user, "address", None),
        avatar_url=getattr(user, "avatar_url", None),
        created_at=user.created_at,
    )


def replace_roles(db: Session, user: User, roles: list[str]) -> None:
    db.query(UserRoleAssignment).filter(
        UserRoleAssignment.user_id == user.id
    ).delete()
    db.add_all(
        UserRoleAssignment(user_id=user.id, role=role)
        for role in roles
    )
    user.role = roles[0]


def replace_class_assignments(
    db: Session, user: User, assigned_classes: list[str]
) -> None:
    db.query(UserClassAssignment).filter(
        UserClassAssignment.user_id == user.id
    ).delete()
    unique_classes = list(dict.fromkeys(name.strip() for name in assigned_classes if name.strip()))
    db.add_all(
        UserClassAssignment(user_id=user.id, class_name=class_name)
        for class_name in unique_classes
    )


def permission_allows(roles: list[str], module: str, action: str) -> bool:
    required_level = {"read": 1, "write": 2, "full": 3}.get(action)
    if required_level is None:
        return False
    return any(
        PERMISSION_LEVELS.get(ROLE_PERMISSIONS.get(role, {}).get(module, "-"), 0)
        >= required_level
        for role in roles
    )


def hash_password(password: str) -> str:
    salt = secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 120_000)
    return f"{salt}${digest.hex()}"


def verify_password(password: str, stored_hash: str) -> bool:
    try:
        salt, expected = stored_hash.split("$", 1)
    except ValueError:
        return False
    actual = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 120_000).hex()
    return hmac.compare_digest(actual, expected)


def encode_token(user: User, session_started_at: int | None = None) -> str:
    now = int(datetime.now(timezone.utc).timestamp())
    session_started_at = session_started_at or now
    absolute_expiry = session_started_at + TOKEN_TTL_HOURS * 60 * 60
    payload = {
        "sub": user.id,
        "ver": user.auth_version,
        "jti": secrets.token_urlsafe(18),
        "started": session_started_at,
        "last": now,
        "max_exp": absolute_expiry,
        "exp": min(now + SESSION_IDLE_MINUTES * 60, absolute_expiry),
    }
    body = base64.urlsafe_b64encode(
        json.dumps(payload, separators=(",", ":")).encode()
    ).decode()
    signature = hmac.new(TOKEN_SECRET.encode(), body.encode(), hashlib.sha256).hexdigest()
    return f"{body}.{signature}"


def decode_token(token: str) -> dict[str, Any]:
    try:
        body, signature = token.split(".", 1)
        expected = hmac.new(TOKEN_SECRET.encode(), body.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(signature, expected):
            raise ValueError
        payload = json.loads(base64.urlsafe_b64decode(body.encode()))
        numeric_claims = ("sub", "ver", "started", "last", "max_exp", "exp")
        if not isinstance(payload, dict) or any(
            not isinstance(payload.get(claim), int)
            or isinstance(payload.get(claim), bool)
            for claim in numeric_claims
        ):
            raise ValueError
        now = int(datetime.now(timezone.utc).timestamp())
        if (
            payload["exp"] <= now
            or payload["last"] + SESSION_IDLE_MINUTES * 60 <= now
            or payload["max_exp"] <= now
        ):
            raise ValueError
        return payload
    except (ValueError, KeyError, json.JSONDecodeError):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Phiên đăng nhập không hợp lệ",
        ) from None


def get_current_user(
    credentials: CurrentCredentials,
    db: DbSession,
) -> User:
    if credentials is None:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED, detail="Vui lòng đăng nhập"
        )
    payload = decode_token(credentials.credentials)
    now = int(datetime.now(timezone.utc).timestamp())
    user = db.get(User, payload["sub"])
    if (
        user is None
        or not user.is_active
        or payload.get("ver") != user.auth_version
        or payload["last"] + SESSION_IDLE_MINUTES * 60 < now
        or payload["max_exp"] < now
    ):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Phiên đăng nhập không còn hiệu lực",
        )
    return user


CurrentUser = Annotated[User, Depends(get_current_user)]


def require_permission(module: str, action: str):
    def permission_gate(user: CurrentUser, db: DbSession) -> User:
        if not permission_allows(assigned_roles(db, user), module, action):
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Bạn không có quyền truy cập chức năng này",
            )
        return user

    return permission_gate


def seed_database() -> None:
    Base.metadata.create_all(engine)
    existing_columns = {column["name"] for column in inspect(engine).get_columns("users")}
    migrations = {
        "failed_login_attempts": "INTEGER NOT NULL DEFAULT 0",
        "locked_until": "DATETIME NULL",
        "phone": "VARCHAR(40) NULL",
        "lock_reason": "VARCHAR(500) NULL",
        "auth_version": "INTEGER NOT NULL DEFAULT 0",
        "must_change_password": "BOOLEAN NOT NULL DEFAULT 0",
        "date_of_birth": "DATE NULL",
        "address": "VARCHAR(500) NULL",
        "avatar_url": "VARCHAR(500) NULL",
    }
    with engine.begin() as connection:
        for column, definition in migrations.items():
            if column not in existing_columns:
                connection.execute(text(f"ALTER TABLE users ADD COLUMN {column} {definition}"))
        for table, additions in {
            "leads": {
                "loss_reason": "VARCHAR(1000) NULL",
                "converted_at": "DATETIME NULL",
            },
            "lead_interactions": {
                "interaction_type": "VARCHAR(20) NOT NULL DEFAULT 'message'",
                "follow_up_at": "DATETIME NULL",
            },
        }.items():
            columns = {column["name"] for column in inspect(engine).get_columns(table)}
            for column, definition in additions.items():
                if column not in columns:
                    connection.execute(
                        text(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")
                    )
        if engine.dialect.name == "sqlite":
            connection.exec_driver_sql(
                "CREATE INDEX IF NOT EXISTS ix_leads_status ON leads(status)"
            )
            connection.exec_driver_sql(
                "CREATE INDEX IF NOT EXISTS ix_lead_interactions_follow_up_at "
                "ON lead_interactions(follow_up_at)"
            )
        connection.execute(
            text(
                "UPDATE leads SET assigned_to_user_id = NULL "
                "WHERE assigned_to_user_id IS NOT NULL "
                "AND NOT EXISTS (SELECT 1 FROM users "
                "WHERE users.id = leads.assigned_to_user_id)"
            )
        )
    with SessionLocal() as db:
        for user in db.scalars(select(User)).all():
            if user.role == "lecturer":
                user.role = "instructor"
            if user.role not in ROLE_CODES:
                user.role = "guest"
            if db.get(UserRoleAssignment, (user.id, user.role)) is None:
                db.add(UserRoleAssignment(user_id=user.id, role=user.role))
        if db.scalar(select(User).limit(1)) is None:
            demo_users = [
                ("Quản trị viên TMS", "admin@tms.local", "Admin@123", "admin"),
                ("Nguyễn Văn Giảng", "lecturer@tms.local", "Lecturer@123", "instructor"),
                ("Kế toán TMS", "accountant@tms.local", "Accountant@123", "accountant"),
            ]
            for full_name, email, password, role in demo_users:
                user = User(
                    full_name=full_name,
                    email=email,
                    password_hash=hash_password(password),
                    role=role,
                )
                db.add(user)
                db.flush()
                db.add(UserRoleAssignment(user_id=user.id, role=role))
        db.commit()


def smtp_configured() -> bool:
    return all((SMTP_HOST, SMTP_USER, SMTP_PASSWORD, SMTP_FROM))


def send_email(message: EmailMessage) -> bool:
    if not smtp_configured():
        return False
    with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=15) as server:
        server.starttls()
        server.login(SMTP_USER, SMTP_PASSWORD)
        server.send_message(message)
    return True


def branded_email(
    *,
    recipient: str,
    subject: str,
    preheader: str,
    heading: str,
    greeting: str,
    body_html: str,
    action_label: str,
    action_url: str,
    security_note: str,
    text_body: str,
) -> EmailMessage:
    sender_address = parseaddr(SMTP_FROM)[1] or SMTP_FROM or "tms-local@localhost"
    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = formataddr(("Hệ thống Quản lý Đào tạo TMS", sender_address))
    message["To"] = recipient
    message.set_content(text_body)
    message.add_alternative(
        f"""\
<!doctype html>
<html lang="vi">
  <body style="margin:0;padding:0;background:#f3f6fb;font-family:Arial,Helvetica,sans-serif;color:#203047;">
    <div style="display:none;max-height:0;overflow:hidden;opacity:0;">{escape(preheader)}</div>
    <table role="presentation" width="100%" cellspacing="0" cellpadding="0" style="background:#f3f6fb;padding:32px 12px;">
      <tr><td align="center">
        <table role="presentation" width="600" cellspacing="0" cellpadding="0" style="width:100%;max-width:600px;background:#ffffff;border:1px solid #e3eaf2;border-radius:12px;overflow:hidden;">
          <tr><td style="background:#1769e0;padding:22px 32px;color:#ffffff;">
            <div style="font-size:12px;letter-spacing:1.2px;text-transform:uppercase;opacity:.9;">TMS · ĐÀO TẠO</div>
            <div style="font-size:21px;font-weight:700;margin-top:7px;">Hệ thống Quản lý Đào tạo</div>
          </td></tr>
          <tr><td style="padding:32px;">
            <h1 style="font-size:23px;line-height:1.3;margin:0 0 20px;color:#152b45;">{escape(heading)}</h1>
            <p style="font-size:15px;line-height:1.65;margin:0 0 14px;">{escape(greeting)}</p>
            {body_html}
            <p style="margin:26px 0 22px;text-align:center;">
              <a href="{escape(action_url, quote=True)}" style="display:inline-block;background:#1769e0;color:#ffffff;text-decoration:none;font-size:15px;font-weight:700;padding:13px 22px;border-radius:7px;">{escape(action_label)}</a>
            </p>
            <p style="font-size:13px;line-height:1.6;color:#53677e;margin:0 0 10px;">{escape(security_note)}</p>
            <p style="font-size:12px;line-height:1.6;color:#738399;margin:18px 0 0;">Nếu nút không hoạt động, sao chép đường dẫn này vào trình duyệt:<br><a href="{escape(action_url, quote=True)}" style="color:#1769e0;word-break:break-all;">{escape(action_url)}</a></p>
          </td></tr>
          <tr><td style="border-top:1px solid #e8edf4;padding:18px 32px;background:#f9fbfe;color:#718096;font-size:12px;line-height:1.6;">
            Email tự động từ <strong>Hệ thống Quản lý Đào tạo TMS</strong>.<br>
            Vui lòng không trả lời email này. Nếu cần hỗ trợ, liên hệ quản trị viên hệ thống.
          </td></tr>
        </table>
      </td></tr>
    </table>
  </body>
</html>
""",
        subtype="html",
    )
    return message


def reset_email(email: str, token: str) -> EmailMessage:
    reset_url = (
        f"{FRONTEND_URL.rstrip('/')}/?reset_token={quote(token)}"
        f"&account={quote(email)}"
    )
    return branded_email(
        recipient=email,
        subject="Yêu cầu đặt lại mật khẩu | TMS",
        preheader="Liên kết đặt lại mật khẩu có hiệu lực trong 30 phút.",
        heading="Đặt lại mật khẩu",
        greeting="Xin chào,",
        body_html=(
            '<p style="font-size:15px;line-height:1.65;margin:0 0 14px;">'
            "Chúng tôi nhận được yêu cầu đặt lại mật khẩu cho tài khoản:</p>"
            f'<p style="font-size:15px;font-weight:700;color:#152b45;margin:0 0 14px;">{escape(email)}</p>'
            '<p style="font-size:15px;line-height:1.65;margin:0;">'
            "Nhấn nút bên dưới để tạo mật khẩu mới. Liên kết sẽ hết hạn sau <strong>30 phút</strong> "
            "và chỉ sử dụng được <strong>một lần</strong>.</p>"
        ),
        action_label="Đặt lại mật khẩu",
        action_url=reset_url,
        security_note=(
            "Nếu bạn không yêu cầu thay đổi mật khẩu, hãy bỏ qua email này. "
            "Mật khẩu hiện tại của bạn sẽ không thay đổi."
        ),
        text_body=(
            "HỆ THỐNG QUẢN LÝ ĐÀO TẠO TMS\n"
            "YÊU CẦU ĐẶT LẠI MẬT KHẨU\n\n"
            "Xin chào,\n"
            "Chúng tôi nhận được yêu cầu đặt lại mật khẩu cho tài khoản "
            f"{email}.\n"
            f"Mở liên kết sau trong vòng 30 phút: {reset_url}\n"
            "Liên kết chỉ sử dụng được một lần.\n\n"
            "Nếu bạn không yêu cầu thay đổi mật khẩu, hãy bỏ qua email này. "
            "Mật khẩu hiện tại của bạn sẽ không thay đổi.\n\n"
            "Email tự động từ Hệ thống Quản lý Đào tạo TMS. "
            "Vui lòng không trả lời email này."
        ),
    )


def activation_email(email: str, full_name: str, temporary_password: str) -> EmailMessage:
    login_url = FRONTEND_URL.rstrip("/")
    return branded_email(
        recipient=email,
        subject="Thông tin tài khoản của bạn | TMS",
        preheader="Tài khoản TMS đã được tạo. Đăng nhập lần đầu bằng mật khẩu tạm thời.",
        heading="Tài khoản của bạn đã sẵn sàng",
        greeting=f"Xin chào {full_name},",
        body_html=(
            '<p style="font-size:15px;line-height:1.65;margin:0 0 14px;">'
            "Tài khoản của bạn trên Hệ thống Quản lý Đào tạo đã được tạo.</p>"
            '<table role="presentation" width="100%" cellspacing="0" cellpadding="0" '
            'style="background:#f4f7fb;border:1px solid #e3eaf2;border-radius:8px;margin:18px 0;">'
            f'<tr><td style="padding:12px 16px;font-size:14px;">Email đăng nhập: <strong>{escape(email)}</strong></td></tr>'
            f'<tr><td style="padding:0 16px 14px;font-size:14px;">Mật khẩu tạm thời: <strong>{escape(temporary_password)}</strong></td></tr>'
            "</table>"
            '<p style="font-size:15px;line-height:1.65;margin:0;">'
            "Vì lý do bảo mật, bạn sẽ được yêu cầu đổi mật khẩu ngay sau lần đăng nhập đầu tiên.</p>"
        ),
        action_label="Đăng nhập hệ thống",
        action_url=login_url,
        security_note=(
            "Mật khẩu tạm thời được gửi riêng cho bạn. Không chia sẻ email này; "
            "hãy đổi mật khẩu ngay sau khi đăng nhập."
        ),
        text_body=(
            "HỆ THỐNG QUẢN LÝ ĐÀO TẠO TMS\n"
            "TÀI KHOẢN CỦA BẠN ĐÃ ĐƯỢC TẠO\n\n"
            f"Xin chào {full_name},\n\n"
            "Tài khoản của bạn trên Hệ thống Quản lý Đào tạo đã được tạo.\n"
            f"Email đăng nhập: {email}\n"
            f"Mật khẩu tạm thời: {temporary_password}\n"
            f"Trang đăng nhập: {login_url}\n\n"
            "Bạn sẽ được yêu cầu đổi mật khẩu ngay sau lần đăng nhập đầu tiên. "
            "Không chia sẻ mật khẩu tạm thời này.\n\n"
            "Nếu bạn không yêu cầu tạo tài khoản này, vui lòng liên hệ quản trị viên.\n\n"
            "Email tự động từ Hệ thống Quản lý Đào tạo TMS. "
            "Vui lòng không trả lời email này."
        ),
    )


app = FastAPI(title="TMS Training Management API", version="1.0.0", on_startup=[seed_database])
app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Retry-After"],
)
app.mount("/uploads", StaticFiles(directory=UPLOADS_DIR.parent), name="uploads")


@app.get("/")
def read_root() -> dict[str, str]:
    return {"message": "TMS Backend is running", "docs": "/docs"}


@app.get("/health")
def health_check() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/auth/roles")
def list_roles(_: CurrentUser) -> list[dict[str, str]]:
    return [{"code": role, "label": ROLE_LABELS[role]} for role in ROLE_CODES]


@app.get("/auth/me", response_model=UserResponse)
def me(user: CurrentUser, db: DbSession) -> UserResponse:
    return user_response(db, user)


class ProfileUpdateRequest(BaseModel):
    full_name: str | None = None
    phone: str | None = None
    date_of_birth: date | None = None
    address: str | None = None


@app.get("/profile")
def get_profile(user: CurrentUser, db: DbSession) -> dict[str, Any]:
    return {
        "id": user.id,
        "full_name": user.full_name,
        "email": user.email,
        "phone": user.phone,
        "date_of_birth": user.date_of_birth.isoformat() if getattr(user, "date_of_birth", None) else None,
        "address": getattr(user, "address", None),
        "avatar_url": getattr(user, "avatar_url", None),
    }


@app.patch("/profile")
def update_profile(
    user: CurrentUser,
    db: DbSession,
    request: ProfileUpdateRequest,
) -> dict[str, Any]:
    if request.full_name is not None:
        if len(request.full_name.strip()) < 2:
            raise HTTPException(status_code=422, detail="Họ tên phải có ít nhất 2 ký tự")
        user.full_name = request.full_name.strip()
    if request.phone is not None:
        normalized = request.phone.strip()
        if normalized and not re.fullmatch(r"^(\+84|84|0)[0-9]{9,10}$", normalized.replace(" ", "")):
            raise HTTPException(status_code=422, detail="Số điện thoại không hợp lệ")
        user.phone = normalized or None
    if "date_of_birth" in request.model_fields_set:
        user.date_of_birth = request.date_of_birth
    if request.address is not None:
        user.address = request.address.strip()[:500] or None
    db.commit()
    db.refresh(user)
    return get_profile(user, db)


@app.post("/profile/avatar")
async def upload_avatar(
    user: CurrentUser,
    db: DbSession,
    file: UploadFile = File(...),
) -> dict[str, Any]:
    if not file.filename:
        raise HTTPException(status_code=400, detail="Vui lòng chọn ảnh đại diện")
    suffix = Path(file.filename).suffix.lower()
    if suffix not in {".png", ".jpg", ".jpeg", ".webp"}:
        raise HTTPException(status_code=400, detail="Chỉ chấp nhận JPG/PNG/WebP")
    contents = await file.read()
    if len(contents) > 2 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Ảnh tối đa 2MB")
    previous_avatar_url = user.avatar_url
    filename = f"avatar_{user.id}_{secrets.token_hex(8)}{suffix}"
    target = UPLOADS_DIR / filename
    target.write_bytes(contents)
    user.avatar_url = f"/uploads/avatars/{filename}"
    db.commit()
    db.refresh(user)
    if previous_avatar_url and previous_avatar_url.startswith("/uploads/avatars/"):
        previous_filename = Path(previous_avatar_url).name
        if previous_filename.startswith(f"avatar_{user.id}_") or previous_filename.startswith(f"avatar_{user.id}."):
            previous_target = UPLOADS_DIR / previous_filename
            if previous_target != target:
                previous_target.unlink(missing_ok=True)
    return {"avatar_url": user.avatar_url}


def inspect_user_import(content: bytes, db: Session) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        workbook = load_workbook(filename=BytesIO(content), data_only=True, read_only=True)
    except (InvalidFileException, BadZipFile, OSError, ValueError) as error:
        raise HTTPException(status_code=400, detail="Tệp Excel không hợp lệ hoặc bị hỏng") from error
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not rows:
        raise HTTPException(status_code=400, detail="Tệp Excel rỗng")
    headers = [str(value).strip().lower() for value in rows[0]]
    header_map = {header: index for index, header in enumerate(headers)}
    required = ["full_name", "email", "roles"]
    missing = [item for item in required if item not in header_map]
    if missing:
        raise HTTPException(status_code=422, detail=f"Thiếu cột bắt buộc: {', '.join(missing)}")
    errors: list[str] = []
    valid_rows: list[dict[str, Any]] = []
    all_rows: list[dict[str, Any]] = []
    seen_emails: set[str] = set()
    for idx, row in enumerate(rows[1:], start=2):
        full_name = str(row[header_map["full_name"]] or "").strip() if "full_name" in header_map and row[header_map["full_name"]] is not None else ""
        email = str(row[header_map["email"]] or "").strip() if "email" in header_map and row[header_map["email"]] is not None else ""
        roles_value = row[header_map["roles"]] if "roles" in header_map else ""
        phone = str(row[header_map["phone"]] or "") if "phone" in header_map and row[header_map["phone"]] is not None else None
        class_value = str(row[header_map["assigned_classes"]] or "") if "assigned_classes" in header_map and row[header_map["assigned_classes"]] is not None else ""
        assigned_classes = [item.strip() for item in class_value.replace(";", ",").split(",") if item.strip()]

        role_tokens = [token.strip() for token in str(roles_value).replace(";", ",").split(",") if token.strip()]
        if not role_tokens:
            role_tokens = ["student"]
        normalized_roles = [token.lower().strip() for token in role_tokens]

        if not full_name or not email:
            err_msg = "Thiếu họ tên hoặc email"
            errors.append(f"Dòng {idx}: {err_msg}")
            all_rows.append({
                "row_number": idx, "full_name": full_name, "email": email,
                "roles": normalized_roles, "phone": phone, "assigned_classes": assigned_classes,
                "status": "error", "message": err_msg,
            })
            continue

        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            err_msg = f"Email không hợp lệ ({email})"
            errors.append(f"Dòng {idx}: email không hợp lệ: {email}")
            all_rows.append({
                "row_number": idx, "full_name": full_name, "email": email,
                "roles": normalized_roles, "phone": phone, "assigned_classes": assigned_classes,
                "status": "error", "message": err_msg,
            })
            continue

        normalized_email = email.lower()
        if normalized_email in seen_emails or db.scalar(select(User).where(User.email == normalized_email)) is not None:
            err_msg = f"Email đã tồn tại ({email})"
            errors.append(f"Dòng {idx}: email đã tồn tại: {email}")
            all_rows.append({
                "row_number": idx, "full_name": full_name, "email": email,
                "roles": normalized_roles, "phone": phone, "assigned_classes": assigned_classes,
                "status": "duplicate", "message": err_msg,
            })
            continue

        invalid_roles = [role for role in normalized_roles if role not in ROLE_CODES]
        if invalid_roles:
            err_msg = f"Vai trò không hợp lệ: {', '.join(invalid_roles)}"
            errors.append(f"Dòng {idx}: {err_msg}")
            all_rows.append({
                "row_number": idx, "full_name": full_name, "email": email,
                "roles": normalized_roles, "phone": phone, "assigned_classes": assigned_classes,
                "status": "error", "message": err_msg,
            })
            continue

        initial_password = str(row[header_map["initial_password"]] or "") if "initial_password" in header_map and row[header_map["initial_password"]] is not None else ""
        if not smtp_configured():
            try:
                validate_password_strength(initial_password)
            except ValueError as error:
                err_msg = f"Cần mật khẩu khởi tạo hợp lệ: {error}"
                errors.append(f"Dòng {idx}: {err_msg}")
                all_rows.append({
                    "row_number": idx, "full_name": full_name, "email": email,
                    "roles": normalized_roles, "phone": phone, "assigned_classes": assigned_classes,
                    "status": "error", "message": err_msg,
                })
                continue

        valid_row_item = {
            "row_number": idx,
            "full_name": full_name,
            "email": normalized_email,
            "roles": list(dict.fromkeys(normalized_roles)),
            "phone": phone,
            "initial_password": initial_password,
            "assigned_classes": assigned_classes,
        }
        valid_rows.append(valid_row_item)
        all_rows.append({
            "row_number": idx,
            "full_name": full_name,
            "email": normalized_email,
            "roles": list(dict.fromkeys(normalized_roles)),
            "phone": phone,
            "assigned_classes": assigned_classes,
            "status": "valid",
            "message": "Hợp lệ",
        })
        seen_emails.add(normalized_email)
    return valid_rows, errors, all_rows


@app.post("/users/import-excel/preview")
def preview_user_import(
    _: Annotated[User, Depends(require_permission("users", "full"))],
    db: DbSession,
    file: UploadFile = File(...),
) -> dict[str, Any]:
    if Path(file.filename or "").suffix.lower() != ".xlsx":
        raise HTTPException(status_code=400, detail="Chỉ chấp nhận tệp Excel .xlsx")
    valid_rows, errors, all_rows = inspect_user_import(file.file.read(), db)
    duplicate_count = sum(1 for r in all_rows if r.get("status") == "duplicate")
    error_count = sum(1 for r in all_rows if r.get("status") == "error")
    preview_rows = [
        {key: value for key, value in row.items() if key != "initial_password"}
        for row in valid_rows
    ]
    return {
        "valid": len(valid_rows),
        "skipped": len(errors),
        "duplicates": duplicate_count,
        "error_count": error_count,
        "errors": errors,
        "preview_rows": preview_rows,
        "all_rows": all_rows,
    }


@app.post("/users/import-excel")
def import_users_excel(
    user: Annotated[User, Depends(require_permission("users", "full"))],
    db: DbSession,
    file: UploadFile = File(...),
) -> dict[str, Any]:
    if Path(file.filename or "").suffix.lower() != ".xlsx":
        raise HTTPException(status_code=400, detail="Chỉ chấp nhận tệp Excel .xlsx")
    valid_rows, errors, _ = inspect_user_import(file.file.read(), db)
    imported = 0
    for row in valid_rows:
        temporary_password = row["initial_password"] or f"{secrets.token_urlsafe(10)}A1!"
        if smtp_configured():
            try:
                email_sent = send_email(activation_email(row["email"], row["full_name"], temporary_password))
            except (smtplib.SMTPException, OSError):
                logger.exception("Unable to send import activation")
                email_sent = False
            if not email_sent:
                errors.append(f"Dòng {row['row_number']}: không gửi được email kích hoạt tới {row['email']}; tài khoản chưa được tạo")
                continue
        new_user = User(
            full_name=row["full_name"],
            email=row["email"],
            password_hash=hash_password(temporary_password),
            role=row["roles"][0],
            must_change_password=True,
            phone=row["phone"],
        )
        db.add(new_user)
        db.flush()
        replace_roles(db, new_user, row["roles"])
        replace_class_assignments(db, new_user, row["assigned_classes"])
        imported += 1
    db.commit()
    return {"imported": imported, "skipped": len(errors), "errors": errors}


@app.get("/users/import-excel/template")
def download_user_import_template(
    _: Annotated[User, Depends(require_permission("users", "full"))],
) -> Response:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Users"
    sheet.append(["full_name", "email", "roles", "phone", "assigned_classes", "initial_password"])
    sheet.append(["Nguyen Van A", "nguyenvana@example.com", "student", "0912345678", "Lop A", ""])
    sheet.freeze_panes = "A2"
    for column, width in {"A": 24, "B": 34, "C": 22, "D": 18, "E": 24, "F": 24}.items():
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="users-template.xlsx"'},
    )


@app.get("/training/programs")
def list_programs(
    user: Annotated[User, Depends(require_permission("courses", "read"))],
    db: DbSession,
    q: str | None = None,
    page: int | None = Query(default=None, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
) -> list[dict[str, Any]] | dict[str, Any]:
    statement = select(TrainingProgram)
    if q and q.strip():
        search = f"%{q.strip()}%"
        statement = statement.where(
            TrainingProgram.code.ilike(search) | TrainingProgram.name.ilike(search)
        )
    total = db.scalar(select(func.count()).select_from(statement.subquery())) or 0
    ordered = statement.order_by(TrainingProgram.created_at.desc(), TrainingProgram.id.desc())
    if page is None:
        programs = db.scalars(ordered).all()
        return [
            {
                "id": program.id,
                "code": program.code,
                "name": program.name,
                "description": program.description,
                "total_duration_hours": program.total_duration_hours,
                "standard_fee": program.standard_fee,
                "status": program.status,
            }
            for program in programs
        ]
    programs = db.scalars(
        ordered.offset((page - 1) * page_size).limit(page_size)
    ).all()
    return {
        "items": [
            {
                "id": program.id,
                "code": program.code,
                "name": program.name,
                "description": program.description,
                "total_duration_hours": program.total_duration_hours,
                "standard_fee": program.standard_fee,
                "status": program.status,
            }
            for program in programs
        ],
        "total": total,
        "page": page,
        "page_size": page_size,
    }


@app.post("/training/programs")
def create_program(
    user: Annotated[User, Depends(require_permission("courses", "full"))],
    db: DbSession,
    payload: dict[str, Any],
) -> dict[str, Any]:
    code = str(payload.get("code", "")).strip()
    name = str(payload.get("name", "")).strip()
    if not code or not name:
        raise HTTPException(status_code=422, detail="Mã và tên chương trình là bắt buộc")
    if len(code) > 30 or len(name) > 200:
        raise HTTPException(status_code=422, detail="Mã chương trình tối đa 30 ký tự, tên tối đa 200 ký tự")
    if db.scalar(select(TrainingProgram).where(func.lower(TrainingProgram.code) == code.lower())):
        raise HTTPException(status_code=409, detail="Mã chương trình đã tồn tại")
    total_duration_hours = int(payload.get("total_duration_hours", 0) or 0)
    standard_fee = int(payload.get("standard_fee", 0) or 0)
    program_status = str(payload.get("status", "active"))
    if total_duration_hours < 0 or standard_fee < 0 or program_status not in {"active", "inactive"}:
        raise HTTPException(status_code=422, detail="Thời lượng, học phí hoặc trạng thái chương trình không hợp lệ")
    program = TrainingProgram(
        code=code,
        name=name,
        description=payload.get("description"),
        total_duration_hours=total_duration_hours,
        standard_fee=standard_fee,
        status=program_status,
    )
    db.add(program)
    db.commit()
    db.refresh(program)
    return {"id": program.id, "code": program.code, "name": program.name, "status": program.status}


@app.patch("/training/programs/{program_id}")
def update_program(
    program_id: int,
    user: Annotated[User, Depends(require_permission("courses", "full"))],
    db: DbSession,
    payload: dict[str, Any],
) -> dict[str, Any]:
    program = db.get(TrainingProgram, program_id)
    if not program:
        raise HTTPException(status_code=404, detail="Không tìm thấy chương trình")
    if "code" in payload and payload["code"]:
        code = str(payload["code"]).strip();
        if len(code) > 30:
            raise HTTPException(status_code=422, detail="Mã chương trình tối đa 30 ký tự")
        if code.lower() != program.code.lower() and db.scalar(select(TrainingProgram).where(func.lower(TrainingProgram.code) == code.lower())):
            raise HTTPException(status_code=409, detail="Mã chương trình đã tồn tại")
        program.code = code
    if "name" in payload and payload["name"]:
        name = str(payload["name"]).strip()
        if len(name) > 200:
            raise HTTPException(status_code=422, detail="Tên chương trình tối đa 200 ký tự")
        program.name = name
    if "description" in payload:
        program.description = payload["description"]
    if "total_duration_hours" in payload:
        duration = int(payload["total_duration_hours"] or 0)
        if duration < 0:
            raise HTTPException(status_code=422, detail="Thời lượng không được âm")
        program.total_duration_hours = duration
    if "standard_fee" in payload:
        fee = int(payload["standard_fee"] or 0)
        if fee < 0:
            raise HTTPException(status_code=422, detail="Học phí không được âm")
        program.standard_fee = fee
    if "status" in payload:
        new_status = str(payload["status"])
        if new_status not in {"active", "inactive"}:
            raise HTTPException(status_code=422, detail="Trạng thái chương trình không hợp lệ")
        program.status = new_status
    db.commit(); db.refresh(program)
    return {"id": program.id, "code": program.code, "name": program.name, "status": program.status}


@app.delete("/training/programs/{program_id}")
def delete_program(
    program_id: int,
    user: Annotated[User, Depends(require_permission("courses", "full"))],
    db: DbSession,
) -> dict[str, str]:
    program = db.get(TrainingProgram, program_id)
    if not program:
        raise HTTPException(status_code=404, detail="Không tìm thấy chương trình")
    running_class = db.scalar(
        select(TrainingClass).where(
            TrainingClass.program_id == program_id,
            TrainingClass.status == "running",
        ).limit(1)
    )
    if running_class:
        raise HTTPException(
            status_code=409,
            detail="Chương trình đang có lớp diễn ra; hãy ngừng áp dụng chương trình thay vì xoá",
        )
    db.execute(delete(TrainingClass).where(TrainingClass.program_id == program_id))
    db.execute(delete(ProgramSubject).where(ProgramSubject.program_id == program_id))
    db.delete(program)
    db.commit()
    return {"message": "Đã xoá chương trình"}


@app.get("/training/programs/{program_id}/classes")
def list_training_classes(
    program_id: int,
    user: Annotated[User, Depends(require_permission("courses", "read"))],
    db: DbSession,
) -> list[dict[str, Any]]:
    if db.get(TrainingProgram, program_id) is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy chương trình")
    classes = db.scalars(
        select(TrainingClass)
        .where(TrainingClass.program_id == program_id)
        .order_by(TrainingClass.created_at.desc())
    ).all()
    return [{"id": item.id, "name": item.name, "status": item.status} for item in classes]


@app.post("/training/programs/{program_id}/classes", status_code=status.HTTP_201_CREATED)
def create_training_class(
    program_id: int,
    user: Annotated[User, Depends(require_permission("courses", "full"))],
    db: DbSession,
    payload: dict[str, Any],
) -> dict[str, Any]:
    program = db.get(TrainingProgram, program_id)
    if program is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy chương trình")
    if program.status != "active":
        raise HTTPException(status_code=409, detail="Không thể mở lớp từ chương trình đã ngừng áp dụng")
    name = str(payload.get("name", "")).strip()
    if not name or len(name) > 200:
        raise HTTPException(status_code=422, detail="Tên lớp là bắt buộc và tối đa 200 ký tự")
    training_class = TrainingClass(program_id=program_id, name=name, status="planned")
    db.add(training_class)
    db.commit()
    db.refresh(training_class)
    return {"id": training_class.id, "name": training_class.name, "status": training_class.status}


@app.patch("/training/classes/{class_id}")
def update_training_class(
    class_id: int,
    user: Annotated[User, Depends(require_permission("courses", "full"))],
    db: DbSession,
    payload: dict[str, Any],
) -> dict[str, Any]:
    training_class = db.get(TrainingClass, class_id)
    if training_class is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lớp học")
    if "name" in payload:
        name = str(payload["name"]).strip()
        if not name or len(name) > 200:
            raise HTTPException(status_code=422, detail="Tên lớp là bắt buộc và tối đa 200 ký tự")
        training_class.name = name
    if "status" in payload:
        new_status = str(payload["status"])
        if new_status not in {"planned", "running", "completed", "cancelled"}:
            raise HTTPException(status_code=422, detail="Trạng thái lớp không hợp lệ")
        program = db.get(TrainingProgram, training_class.program_id)
        if new_status == "running" and program and program.status != "active":
            raise HTTPException(status_code=409, detail="Không thể chạy lớp thuộc chương trình đã ngừng áp dụng")
        training_class.status = new_status
    db.commit()
    db.refresh(training_class)
    return {"id": training_class.id, "name": training_class.name, "status": training_class.status}


@app.get("/training/subjects")
def list_subjects(
    user: Annotated[User, Depends(require_permission("courses", "read"))],
    db: DbSession,
    q: str | None = None,
    page: int | None = Query(default=None, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
) -> list[dict[str, Any]] | dict[str, Any]:
    statement = select(TrainingSubject)
    if q and q.strip():
        search = f"%{q.strip()}%"
        statement = statement.where(
            TrainingSubject.code.ilike(search) | TrainingSubject.name.ilike(search)
        )
    total = db.scalar(select(func.count()).select_from(statement.subquery())) or 0
    ordered = statement.order_by(TrainingSubject.name, TrainingSubject.id)
    if page is None:
        items = db.scalars(ordered).all()
        return [
            {
                "id": item.id,
                "code": item.code,
                "name": item.name,
                "session_count": item.session_count,
                "weight": item.weight,
                "description": item.description,
                "learning_outcomes": item.learning_outcomes,
            }
            for item in items
        ]
    items = db.scalars(
        ordered.offset((page - 1) * page_size).limit(page_size)
    ).all()
    return {
        "items": [
            {
                "id": item.id,
                "code": item.code,
                "name": item.name,
                "session_count": item.session_count,
                "weight": item.weight,
                "description": item.description,
                "learning_outcomes": item.learning_outcomes,
            }
            for item in items
        ],
        "total": total,
        "page": page,
        "page_size": page_size,
    }


@app.post("/training/subjects")
def create_subject(
    user: Annotated[User, Depends(require_permission("courses", "full"))],
    db: DbSession,
    payload: dict[str, Any],
) -> dict[str, Any]:
    code = str(payload.get("code", "")).strip(); name = str(payload.get("name", "")).strip()
    if not code or not name:
        raise HTTPException(status_code=422, detail="Mã và tên môn học là bắt buộc")
    if len(code) > 30 or len(name) > 200:
        raise HTTPException(status_code=422, detail="Mã môn tối đa 30 ký tự, tên tối đa 200 ký tự")
    if db.scalar(select(TrainingSubject).where(func.lower(TrainingSubject.code) == code.lower())):
        raise HTTPException(status_code=409, detail="Mã môn học đã tồn tại")
    session_count = int(payload.get("session_count", 0) or 0)
    weight = int(payload.get("weight", 1) or 1)
    if session_count < 0 or weight < 1:
        raise HTTPException(status_code=422, detail="Số buổi không được âm và trọng số phải lớn hơn 0")
    subject = TrainingSubject(
        code=code,
        name=name,
        session_count=session_count,
        weight=weight,
        description=payload.get("description"),
        learning_outcomes=payload.get("learning_outcomes"),
    )
    db.add(subject); db.commit(); db.refresh(subject)
    return {"id": subject.id, "code": subject.code, "name": subject.name}


@app.patch("/training/subjects/{subject_id}")
def update_subject(
    subject_id: int,
    user: Annotated[User, Depends(require_permission("courses", "full"))],
    db: DbSession,
    payload: dict[str, Any],
) -> dict[str, Any]:
    subject = db.get(TrainingSubject, subject_id)
    if not subject:
        raise HTTPException(status_code=404, detail="Không tìm thấy môn học")
    if "code" in payload and payload["code"]:
        code = str(payload["code"]).strip();
        if len(code) > 30:
            raise HTTPException(status_code=422, detail="Mã môn tối đa 30 ký tự")
        if code.lower() != subject.code.lower() and db.scalar(select(TrainingSubject).where(func.lower(TrainingSubject.code) == code.lower())):
            raise HTTPException(status_code=409, detail="Mã môn học đã tồn tại")
        subject.code = code
    if "name" in payload and payload["name"]:
        name = str(payload["name"]).strip()
        if len(name) > 200:
            raise HTTPException(status_code=422, detail="Tên môn tối đa 200 ký tự")
        subject.name = name
    if "session_count" in payload:
        session_count = int(payload["session_count"] or 0)
        if session_count < 0:
            raise HTTPException(status_code=422, detail="Số buổi không được âm")
        max_session_sequence = db.scalar(
            select(func.max(SubjectSession.sequence)).where(SubjectSession.subject_id == subject_id)
        ) or 0
        if session_count > 0 and max_session_sequence > session_count:
            raise HTTPException(status_code=409, detail="Số buổi khai báo không được ít hơn số buổi đã tạo")
        subject.session_count = session_count
    if "weight" in payload:
        weight = int(payload["weight"] or 1)
        if weight < 1:
            raise HTTPException(status_code=422, detail="Trọng số phải lớn hơn 0")
        subject.weight = weight
    if "description" in payload: subject.description = payload["description"]
    if "learning_outcomes" in payload: subject.learning_outcomes = payload["learning_outcomes"]
    db.commit(); db.refresh(subject)
    return {"id": subject.id, "code": subject.code, "name": subject.name}


@app.delete("/training/subjects/{subject_id}")
def delete_subject(
    subject_id: int,
    user: Annotated[User, Depends(require_permission("courses", "full"))],
    db: DbSession,
) -> dict[str, str]:
    subject = db.get(TrainingSubject, subject_id)
    if not subject:
        raise HTTPException(status_code=404, detail="Không tìm thấy môn học")
    running_use = db.scalar(
        select(TrainingClass.id)
        .join(ProgramSubject, ProgramSubject.program_id == TrainingClass.program_id)
        .where(
            ProgramSubject.subject_id == subject_id,
            TrainingClass.status == "running",
        )
        .limit(1)
    )
    if running_use:
        raise HTTPException(status_code=409, detail="Môn học đang được sử dụng trong lớp diễn ra; không thể xoá")
    db.execute(delete(ProgramSubject).where(ProgramSubject.subject_id == subject_id))
    db.execute(
        update(ProgramSubject)
        .where(ProgramSubject.prerequisite_subject_id == subject_id)
        .values(prerequisite_subject_id=None)
    )
    db.execute(delete(SubjectSession).where(SubjectSession.subject_id == subject_id))
    db.delete(subject)
    db.commit()
    return {"message": "Đã xoá môn học"}


@app.get("/training/programs/{program_id}/subjects")
def list_program_subjects(
    program_id: int,
    user: Annotated[User, Depends(require_permission("courses", "read"))],
    db: DbSession,
) -> list[dict[str, Any]]:
    if db.get(TrainingProgram, program_id) is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy chương trình")
    rows = db.scalars(select(ProgramSubject).where(ProgramSubject.program_id == program_id).order_by(ProgramSubject.sequence_order.asc())).all()
    payload = []
    for row in rows:
        subject = db.get(TrainingSubject, row.subject_id)
        payload.append({
            "id": row.id,
            "subject_id": row.subject_id,
            "code": subject.code if subject else "",
            "name": subject.name if subject else "",
            "sequence_order": row.sequence_order,
            "prerequisite_subject_id": row.prerequisite_subject_id,
            "prerequisite_name": (
                db.get(TrainingSubject, row.prerequisite_subject_id).name
                if row.prerequisite_subject_id and db.get(TrainingSubject, row.prerequisite_subject_id)
                else None
            ),
            "is_required": row.is_required,
        })
    return payload


@app.post("/training/programs/{program_id}/subjects")
def add_subject_to_program(
    program_id: int,
    user: Annotated[User, Depends(require_permission("courses", "full"))],
    db: DbSession,
    payload: dict[str, Any],
) -> dict[str, Any]:
    program = db.get(TrainingProgram, program_id)
    if not program: raise HTTPException(404, "Không tìm thấy chương trình")
    try:
        subject_id = int(payload.get("subject_id"))
        sequence_order = int(payload.get("sequence_order", 1) or 1)
        prerequisite_id = int(payload["prerequisite_subject_id"]) if payload.get("prerequisite_subject_id") else None
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="Môn học, thứ tự hoặc môn tiên quyết không hợp lệ")
    subject = db.get(TrainingSubject, subject_id)
    if not subject: raise HTTPException(404, "Không tìm thấy môn học")
    if sequence_order < 1:
        raise HTTPException(status_code=422, detail="Thứ tự môn học phải từ 1 trở lên")
    if db.scalar(select(ProgramSubject).where(
        ProgramSubject.program_id == program_id,
        ProgramSubject.subject_id == subject.id,
    )):
        raise HTTPException(status_code=409, detail="Môn học đã được gán vào chương trình")
    if prerequisite_id is not None:
        if prerequisite_id == subject.id:
            raise HTTPException(status_code=422, detail="Môn học không thể là tiên quyết của chính nó")
        prerequisite = db.scalar(select(ProgramSubject).where(
            ProgramSubject.program_id == program_id,
            ProgramSubject.subject_id == prerequisite_id,
        ))
        if prerequisite is None:
            raise HTTPException(status_code=422, detail="Môn tiên quyết phải thuộc cùng chương trình")
    existing = db.scalars(
        select(ProgramSubject)
        .where(ProgramSubject.program_id == program_id)
        .order_by(ProgramSubject.sequence_order.asc(), ProgramSubject.id.asc())
    ).all()
    insert_at = min(sequence_order - 1, len(existing))
    association = ProgramSubject(
        program_id=program.id,
        subject_id=subject.id,
        sequence_order=insert_at + 1,
        prerequisite_subject_id=prerequisite_id,
        is_required=bool(payload.get("is_required", True)),
    )
    db.add(association)
    db.flush()
    ordered = [*existing[:insert_at], association, *existing[insert_at:]]
    for order, row in enumerate(ordered, start=1):
        row.sequence_order = order
    db.commit(); db.refresh(association)
    return {"id": association.id, "subject_id": subject.id, "code": subject.code, "name": subject.name}


@app.delete("/training/programs/{program_id}/subjects/{association_id}")
def remove_subject_from_program(
    program_id: int,
    association_id: int,
    user: Annotated[User, Depends(require_permission("courses", "full"))],
    db: DbSession,
) -> dict[str, str]:
    association = db.scalar(select(ProgramSubject).where(
        ProgramSubject.id == association_id,
        ProgramSubject.program_id == program_id,
    ))
    if association is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy môn học trong chương trình")
    db.execute(
        update(ProgramSubject)
        .where(
            ProgramSubject.program_id == program_id,
            ProgramSubject.prerequisite_subject_id == association.subject_id,
        )
        .values(prerequisite_subject_id=None)
    )
    db.delete(association)
    db.commit()
    return {"message": "Đã gỡ môn học khỏi chương trình"}


@app.put("/training/programs/{program_id}/subjects/order")
def reorder_program_subjects(
    program_id: int,
    user: Annotated[User, Depends(require_permission("courses", "full"))],
    db: DbSession,
    payload: dict[str, Any],
) -> list[dict[str, Any]]:
    if db.get(TrainingProgram, program_id) is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy chương trình")
    try:
        ordered_ids = [int(item) for item in payload.get("association_ids", [])]
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="Danh sách thứ tự môn học không hợp lệ")
    associations = db.scalars(select(ProgramSubject).where(ProgramSubject.program_id == program_id)).all()
    existing_ids = {row.id for row in associations}
    if len(ordered_ids) != len(set(ordered_ids)) or set(ordered_ids) != existing_ids:
        raise HTTPException(status_code=422, detail="Danh sách phải gồm chính xác các môn đang thuộc chương trình")
    rows_by_id = {row.id: row for row in associations}
    for order, association_id in enumerate(ordered_ids, start=1):
        rows_by_id[association_id].sequence_order = order
    db.commit()
    return [
        {"id": rows_by_id[association_id].id, "sequence_order": order}
        for order, association_id in enumerate(ordered_ids, start=1)
    ]


@app.get("/training/subjects/{subject_id}/sessions")
def list_subject_sessions(
    subject_id: int,
    user: Annotated[User, Depends(require_permission("courses", "read"))],
    db: DbSession,
) -> list[dict[str, Any]]:
    subject = db.get(TrainingSubject, subject_id)
    if subject is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy môn học")
    rows = db.scalars(select(SubjectSession).where(SubjectSession.subject_id == subject_id).order_by(SubjectSession.sequence.asc())).all()
    if not rows and subject.session_count > 0:
        for seq in range(1, subject.session_count + 1):
            db.add(SubjectSession(
                subject_id=subject.id,
                sequence=seq,
                title=f"Buổi {seq}: Bài học môn {subject.name}",
                objectives=f"Nội dung và mục tiêu kiến thức của buổi {seq} môn {subject.name}."
            ))
        db.commit()
        rows = db.scalars(select(SubjectSession).where(SubjectSession.subject_id == subject_id).order_by(SubjectSession.sequence.asc())).all()
    return [{"id": item.id, "sequence": item.sequence, "title": item.title, "objectives": item.objectives} for item in rows]


@app.post("/training/subjects/{subject_id}/sessions")
def create_subject_session(
    subject_id: int,
    user: Annotated[User, Depends(require_permission("courses", "full"))],
    db: DbSession,
    payload: dict[str, Any],
) -> dict[str, Any]:
    subject = db.get(TrainingSubject, subject_id)
    if not subject: raise HTTPException(404, "Không tìm thấy môn học")
    try:
        sequence = int(payload.get("sequence", 1) or 1)
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="Thứ tự buổi học không hợp lệ")
    session = SubjectSession(
        subject_id=subject.id,
        sequence=sequence,
        title=str(payload.get("title", "")).strip(),
        objectives=payload.get("objectives"),
    )
    if not session.title: raise HTTPException(422, "Tiêu đề buổi học là bắt buộc")
    if session.sequence < 1:
        raise HTTPException(422, "Thứ tự buổi học phải từ 1 trở lên")
    if session.sequence > subject.session_count and subject.session_count > 0:
        raise HTTPException(422, "Số buổi khai báo không vượt quá số buổi của môn học")
    if db.scalar(select(SubjectSession).where(
        SubjectSession.subject_id == subject.id,
        SubjectSession.sequence == session.sequence,
    )):
        raise HTTPException(status_code=409, detail="Thứ tự buổi học đã được sử dụng")
    db.add(session); db.commit(); db.refresh(session)
    return {"id": session.id, "sequence": session.sequence, "title": session.title}


@app.post("/training/subjects/{subject_id}/sessions/clone")
def clone_subject_sessions(
    subject_id: int,
    user: Annotated[User, Depends(require_permission("courses", "full"))],
    db: DbSession,
    payload: dict[str, Any],
) -> dict[str, int]:
    target = db.get(TrainingSubject, subject_id)
    if target is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy môn đích")
    try:
        source_id = int(payload.get("source_subject_id"))
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="Vui lòng chọn môn nguồn")
    source = db.get(TrainingSubject, source_id)
    if source is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy môn nguồn")
    if source.id == target.id:
        raise HTTPException(status_code=422, detail="Môn nguồn và môn đích phải khác nhau")
    source_sessions = db.scalars(
        select(SubjectSession)
        .where(SubjectSession.subject_id == source.id)
        .order_by(SubjectSession.sequence.asc())
    ).all()
    target_sessions = db.scalars(
        select(SubjectSession)
        .where(SubjectSession.subject_id == target.id)
        .order_by(SubjectSession.sequence.asc())
    ).all()
    if not source_sessions:
        raise HTTPException(status_code=422, detail="Môn nguồn chưa có buổi học để sao chép")
    next_sequence = max((item.sequence for item in target_sessions), default=0) + 1
    last_sequence = next_sequence + len(source_sessions) - 1
    if target.session_count > 0 and last_sequence > target.session_count:
        raise HTTPException(status_code=409, detail="Không đủ số buổi cho phép để sao chép toàn bộ buổi học")
    db.add_all([
        SubjectSession(
            subject_id=target.id,
            sequence=next_sequence + offset,
            title=item.title,
            objectives=item.objectives,
        )
        for offset, item in enumerate(source_sessions)
    ])
    db.commit()
    return {"copied_count": len(source_sessions)}


LEAD_PIPELINE_STAGES = {"new", "contacted", "consulting", "trial", "won", "rejected"}
LEGACY_LEAD_STATUSES = {"assigned", "qualified", "converted", "lost"}
LEAD_STATUSES = LEAD_PIPELINE_STAGES | LEGACY_LEAD_STATUSES
LEAD_SOURCES = {"website", "facebook", "walkin", "referral"}


def normalized_lead_phone(phone: str) -> str:
    normalized = phone.replace(" ", "")
    if normalized.startswith("+84"):
        return "0" + normalized[3:]
    if normalized.startswith("84"):
        return "0" + normalized[2:]
    return normalized


def find_duplicate_lead(db: Session, phone: str, exclude_id: int | None = None) -> Lead | None:
    statement = select(Lead)
    if exclude_id is not None:
        statement = statement.where(Lead.id != exclude_id)
    normalized = normalized_lead_phone(phone)
    return next((lead for lead in db.scalars(statement) if normalized_lead_phone(lead.phone) == normalized), None)


def record_lead_interaction(
    db: Session, lead: Lead, user: User | None, action: str, note: str,
    interaction_type: str = "message",
    follow_up_at: datetime | None = None,
    created_at: datetime | None = None,
) -> None:
    db.add(
        LeadInteraction(
            lead_id=lead.id,
            user_id=user.id if user else None,
            action=action,
            note=note[:2000],
            interaction_type=interaction_type,
            follow_up_at=follow_up_at,
            created_at=created_at or datetime.now(timezone.utc),
        )
    )


def ensure_lead_access(db: Session, user: User, lead: Lead) -> None:
    if not permission_allows(assigned_roles(db, user), "leads", "full") and lead.assigned_to_user_id != user.id:
        raise HTTPException(status_code=404, detail="Không tìm thấy lead")


def validate_lead_fields(payload: dict[str, Any], current: Lead | None = None) -> dict[str, Any]:
    full_name = str(payload.get("full_name", current.full_name if current else "")).strip()
    phone = str(payload.get("phone", current.phone if current else "")).strip()
    email = str(payload.get("email", current.email if current else "") or "").strip()
    source = str(payload.get("source", current.source if current else "walkin"))
    program_interest = str(payload.get("program_interest", current.program_interest if current else "") or "").strip()
    notes = str(payload.get("notes", current.notes if current else "") or "").strip()
    if not full_name or not phone:
        raise HTTPException(status_code=422, detail="Họ tên và số điện thoại là bắt buộc")
    if len(full_name) > 120 or len(phone) > 40 or len(program_interest) > 200 or len(notes) > 4000:
        raise HTTPException(status_code=422, detail="Thông tin lead vượt quá độ dài cho phép")
    if not re.fullmatch(r"^(\+84|84|0)[0-9]{9,10}$", phone.replace(" ", "")):
        raise HTTPException(status_code=422, detail="Số điện thoại không hợp lệ")
    if email and (len(email) > 254 or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email)):
        raise HTTPException(status_code=422, detail="Email không hợp lệ")
    if source not in LEAD_SOURCES:
        raise HTTPException(status_code=422, detail="Nguồn lead không hợp lệ")
    return {
        "full_name": full_name,
        "phone": phone.replace(" ", ""),
        "email": email or None,
        "source": source,
        "program_interest": program_interest or None,
        "notes": notes or None,
    }


@app.get("/leads")
def list_leads(
    user: Annotated[User, Depends(require_permission("leads", "read"))],
    db: DbSession,
    q: str | None = None,
    status: str | None = None,
    source: str | None = None,
    assigned_to: int | None = None,
    from_date: date | None = None,
    to_date: date | None = None,
    page: int | None = Query(default=None, ge=1),
    page_size: int = Query(default=20, ge=1, le=500),
) -> list[dict[str, Any]] | dict[str, Any]:
    if status and status not in LEAD_STATUSES:
        raise HTTPException(status_code=422, detail="Trạng thái lead không hợp lệ")
    if source and source not in LEAD_SOURCES:
        raise HTTPException(status_code=422, detail="Nguồn lead không hợp lệ")
    if from_date and to_date and from_date > to_date:
        raise HTTPException(status_code=422, detail="Ngày bắt đầu không được sau ngày kết thúc")
    statement = select(Lead)
    if q:
        q_like = f"%{q.strip()}%"
        statement = statement.where((Lead.full_name.ilike(q_like)) | (Lead.phone.ilike(q_like)) | (Lead.email.ilike(q_like)))
    if status:
        statement = statement.where(Lead.status == status)
    if source:
        statement = statement.where(Lead.source == source)
    if assigned_to is not None:
        statement = statement.where(Lead.assigned_to_user_id == assigned_to)
    if from_date:
        statement = statement.where(Lead.created_at >= datetime.combine(from_date, datetime.min.time()))
    if to_date:
        statement = statement.where(Lead.created_at < datetime.combine(to_date + timedelta(days=1), datetime.min.time()))
    if not permission_allows(assigned_roles(db, user), "leads", "full"):
        statement = statement.where(Lead.assigned_to_user_id == user.id)
    ordered_statement = statement.order_by(Lead.created_at.desc(), Lead.id.desc())
    if page is not None:
        total = db.scalar(select(func.count()).select_from(statement.subquery())) or 0
        ordered_statement = ordered_statement.offset((page - 1) * page_size).limit(page_size)
    leads = db.scalars(ordered_statement).all()
    result = []
    for item in leads:
        assigned_user = (
            db.get(User, item.assigned_to_user_id)
            if item.assigned_to_user_id is not None
            else None
        )
        result.append({
            "id": item.id,
            "full_name": item.full_name,
            "phone": item.phone,
            "email": item.email,
            "source": item.source,
            "status": item.status,
            "program_interest": item.program_interest,
            "notes": item.notes,
            "assigned_to_user_id": assigned_user.id if assigned_user else None,
            "assigned_to_name": assigned_user.full_name if assigned_user else None,
            "created_at": item.created_at.isoformat(),
        })
    if page is not None:
        return {"items": result, "total": total, "page": page, "page_size": page_size}
    return result


@app.get("/leads/duplicate-check")
def check_duplicate_lead(
    _: Annotated[User, Depends(require_permission("leads", "write"))],
    db: DbSession,
    phone: str,
    exclude_id: int | None = None,
) -> dict[str, bool]:
    if not re.fullmatch(r"^(\+84|84|0)[0-9]{9,10}$", phone.replace(" ", "")):
        raise HTTPException(status_code=422, detail="Số điện thoại không hợp lệ")
    return {"duplicate": find_duplicate_lead(db, phone, exclude_id) is not None}


@app.post("/leads", status_code=status.HTTP_201_CREATED)
def create_lead_internal(
    user: Annotated[User, Depends(require_permission("leads", "write"))],
    db: DbSession,
    payload: dict[str, Any],
) -> dict[str, Any]:
    fields = validate_lead_fields(payload)
    if find_duplicate_lead(db, fields["phone"]):
        raise HTTPException(status_code=409, detail="Số điện thoại đã tồn tại trong lead khác")
    lead = Lead(**fields, status="new", assigned_to_user_id=user.id)
    db.add(lead)
    db.flush()
    record_lead_interaction(db, lead, user, "created", "Tạo lead")
    if fields["notes"]:
        record_lead_interaction(db, lead, user, "note", fields["notes"])
    db.commit()
    db.refresh(lead)
    return {"id": lead.id, "full_name": lead.full_name, "status": lead.status}


@app.get("/leads/assignees")
def list_lead_assignees(
    _: Annotated[User, Depends(require_permission("leads", "write"))],
    db: DbSession,
) -> list[dict[str, Any]]:
    candidates = db.scalars(select(User).where(User.is_active.is_(True)).order_by(User.full_name)).all()
    return [
        {"id": candidate.id, "full_name": candidate.full_name}
        for candidate in candidates
        if permission_allows(assigned_roles(db, candidate), "leads", "write")
    ]


@app.post("/leads/public")
def create_lead_public(
    request: Request,
    payload: dict[str, Any],
    db: DbSession,
) -> dict[str, str]:
    if payload.get("website"):
        raise HTTPException(status_code=400, detail="Yêu cầu không hợp lệ")
    fields = validate_lead_fields({**payload, "source": "website"})
    if len(fields["notes"] or "") > 2000:
        raise HTTPException(status_code=422, detail="Ghi chú tối đa 2000 ký tự")
    if find_duplicate_lead(db, fields["phone"]):
        raise HTTPException(status_code=409, detail="Số điện thoại đã được ghi nhận trước đó")
    client_ip = request.client.host if request.client else "unknown"
    client_hash = hmac.new(TOKEN_SECRET.encode(), client_ip.encode(), hashlib.sha256).hexdigest()
    now = datetime.now(timezone.utc)
    window_start = now - timedelta(minutes=10)
    attempts = db.scalar(
        select(func.count(PublicLeadAttempt.id)).where(
            PublicLeadAttempt.client_hash == client_hash,
            PublicLeadAttempt.created_at >= window_start,
        )
    ) or 0
    if attempts >= 5:
        raise HTTPException(
            status_code=429,
            detail="Bạn đã gửi quá nhiều yêu cầu. Vui lòng thử lại sau 10 phút.",
            headers={"Retry-After": "600"},
        )
    db.execute(
        delete(PublicLeadAttempt).where(
            PublicLeadAttempt.created_at < now - timedelta(days=1)
        )
    )
    lead = Lead(**fields, status="new")
    db.add(PublicLeadAttempt(client_hash=client_hash, created_at=now))
    db.add(lead)
    db.flush()
    record_lead_interaction(db, lead, None, "created", "Lead gửi từ biểu mẫu công khai")
    db.commit()
    return {"message": "Đăng ký tư vấn thành công. Bộ phận tuyển sinh sẽ sớm liên hệ với bạn."}


@app.patch("/leads/{lead_id}")
def update_lead(
    lead_id: int,
    user: Annotated[User, Depends(require_permission("leads", "write"))],
    db: DbSession,
    payload: dict[str, Any],
) -> dict[str, Any]:
    lead = db.get(Lead, lead_id)
    if lead is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lead")
    ensure_lead_access(db, user, lead)
    if lead.converted_at is not None:
        raise HTTPException(status_code=409, detail="Lead đã chuyển đổi và đã được khóa")
    editable_fields = {"full_name", "phone", "email", "source", "program_interest", "notes", "status", "loss_reason"}
    if not payload or not set(payload).issubset(editable_fields):
        raise HTTPException(status_code=422, detail="Trường cập nhật lead không hợp lệ")
    changed: list[str] = []
    if set(payload) & {"full_name", "phone", "email", "source", "program_interest", "notes"}:
        fields = validate_lead_fields(payload, lead)
        if normalized_lead_phone(fields["phone"]) != normalized_lead_phone(lead.phone):
            if find_duplicate_lead(db, fields["phone"], lead.id):
                raise HTTPException(status_code=409, detail="Số điện thoại đã tồn tại trong lead khác")
        for key, value in fields.items():
            if key in payload and getattr(lead, key) != value:
                setattr(lead, key, value)
                changed.append(key)
    if "loss_reason" in payload:
        reason = str(payload["loss_reason"] or "").strip()
        if len(reason) > 1000:
            raise HTTPException(status_code=422, detail="Lý do từ chối tối đa 1000 ký tự")
        if lead.loss_reason != (reason or None):
            lead.loss_reason = reason or None
            changed.append("loss_reason")
    if "status" in payload:
        new_status = str(payload["status"])
        if new_status not in LEAD_PIPELINE_STAGES:
            raise HTTPException(status_code=422, detail="Trạng thái lead phải là new, contacted, consulting, trial, won hoặc rejected")
        if new_status == "rejected" and not (
            str(payload.get("loss_reason") or lead.loss_reason or "").strip()
        ):
            raise HTTPException(status_code=422, detail="Khi từ chối lead cần nhập lý do")
        if new_status != lead.status:
            previous_status = lead.status
            lead.status = new_status
            changed.append("status")
            record_lead_interaction(
                db, lead, user, "status_transition",
                f"Chuyển trạng thái: {previous_status} → {new_status}",
            )
    updated_fields = [field for field in changed if field != "status"]
    if updated_fields:
        record_lead_interaction(db, lead, user, "updated", f"Đã cập nhật: {', '.join(updated_fields)}")
    db.commit()
    db.refresh(lead)
    return {"id": lead.id, "status": lead.status, "assigned_to_user_id": lead.assigned_to_user_id}


@app.post("/leads/{lead_id}/assign")
def assign_lead(
    lead_id: int,
    user: Annotated[User, Depends(require_permission("leads", "full"))],
    db: DbSession,
    payload: dict[str, Any],
) -> dict[str, Any]:
    lead = db.get(Lead, lead_id)
    if lead is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lead")
    if lead.converted_at is not None:
        raise HTTPException(status_code=409, detail="Lead đã chuyển đổi và đã được khóa")
    try:
        assigned_to = int(payload.get("assigned_to_user_id") or 0) or None
    except (TypeError, ValueError) as error:
        raise HTTPException(status_code=422, detail="Người phụ trách không hợp lệ") from error
    assignee = db.get(User, assigned_to) if assigned_to else None
    if assigned_to and (
        assignee is None
        or not assignee.is_active
        or not permission_allows(assigned_roles(db, assignee), "leads", "write")
    ):
        raise HTTPException(status_code=404, detail="Tư vấn viên không tồn tại hoặc không hoạt động")
    lead.assigned_to_user_id = assigned_to
    if not assigned_to and lead.status == "assigned":
        lead.status = "new"
    record_lead_interaction(
        db,
        lead,
        user,
        "assignment",
        f"Phân công cho {assignee.full_name}" if assignee else "Gỡ phân công",
    )
    db.commit(); db.refresh(lead)
    return {"id": lead.id, "assigned_to_user_id": lead.assigned_to_user_id, "status": lead.status}


@app.post("/leads/assign-bulk")
def assign_leads_bulk(
    user: Annotated[User, Depends(require_permission("leads", "full"))],
    db: DbSession,
    payload: dict[str, Any],
) -> dict[str, int]:
    raw_ids = payload.get("lead_ids")
    if not isinstance(raw_ids, list) or not raw_ids or len(raw_ids) > 500:
        raise HTTPException(status_code=422, detail="Chọn từ 1 đến 500 lead để phân công")
    try:
        lead_ids = list(dict.fromkeys(int(lead_id) for lead_id in raw_ids))
        assigned_to = int(payload["assigned_to_user_id"]) if payload.get("assigned_to_user_id") else None
    except (TypeError, ValueError, KeyError) as error:
        raise HTTPException(status_code=422, detail="Thông tin phân công không hợp lệ") from error
    leads = db.scalars(select(Lead).where(Lead.id.in_(lead_ids))).all()
    if len(leads) != len(lead_ids):
        raise HTTPException(status_code=404, detail="Một hoặc nhiều lead không tồn tại")
    if any(lead.converted_at is not None for lead in leads):
        raise HTTPException(status_code=409, detail="Lead đã chuyển đổi và đã được khóa")
    assignee = db.get(User, assigned_to) if assigned_to is not None else None
    if assigned_to is not None and (
        assignee is None
        or not assignee.is_active
        or not permission_allows(assigned_roles(db, assignee), "leads", "write")
    ):
        raise HTTPException(status_code=422, detail="Người phụ trách phải là tài khoản tư vấn đang hoạt động")
    for lead in leads:
        lead.assigned_to_user_id = assigned_to
        if not assigned_to and lead.status == "assigned":
            lead.status = "new"
        record_lead_interaction(
            db,
            lead,
            user,
            "assignment",
            f"Phân công hàng loạt cho {assignee.full_name}" if assignee else "Gỡ phân công hàng loạt",
        )
    db.commit()
    return {"assigned_count": len(leads)}


@app.delete("/leads/{lead_id}")
def delete_lead(
    lead_id: int,
    user: Annotated[User, Depends(require_permission("leads", "full"))],
    db: DbSession,
) -> dict[str, str]:
    if "training_manager" not in assigned_roles(db, user):
        raise HTTPException(status_code=403, detail="Chỉ Quản lý đào tạo được xóa lead")
    lead = db.get(Lead, lead_id)
    if lead is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lead")
    if lead.converted_at is not None:
        raise HTTPException(status_code=409, detail="Lead đã chuyển đổi và đã được khóa")
    db.execute(delete(LeadInteraction).where(LeadInteraction.lead_id == lead_id))
    db.delete(lead)
    db.commit()
    return {"message": "Đã xóa lead"}


@app.get("/leads/{lead_id}/interactions")
def list_lead_interactions(
    lead_id: int,
    user: Annotated[User, Depends(require_permission("leads", "read"))],
    db: DbSession,
) -> list[dict[str, Any]]:
    lead = db.get(Lead, lead_id)
    if lead is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lead")
    ensure_lead_access(db, user, lead)
    interactions = db.scalars(
        select(LeadInteraction)
        .where(LeadInteraction.lead_id == lead_id)
        .order_by(LeadInteraction.id.desc())
        .limit(1000)
    ).all()
    return [{
        "id": item.id,
        "action": item.action,
        "note": item.note,
        "interaction_type": item.interaction_type,
        "content": item.note,
        "follow_up_at": (
            (item.follow_up_at.replace(tzinfo=timezone.utc) if item.follow_up_at.tzinfo is None else item.follow_up_at.astimezone(timezone.utc)).isoformat().replace("+00:00", "Z")
            if item.follow_up_at else None
        ),
        "user_name": (
            db.get(User, item.user_id).full_name
            if item.user_id and db.get(User, item.user_id)
            else "Biểu mẫu công khai"
        ),
        "created_at": (
            item.created_at.replace(tzinfo=timezone.utc)
            if item.created_at.tzinfo is None
            else item.created_at.astimezone(timezone.utc)
        ).isoformat().replace("+00:00", "Z"),
    } for item in interactions]


@app.post("/leads/{lead_id}/interactions", status_code=status.HTTP_201_CREATED)
def create_lead_interaction(
    lead_id: int,
    user: Annotated[User, Depends(require_permission("leads", "write"))],
    db: DbSession,
    payload: dict[str, Any],
) -> dict[str, Any]:
    lead = db.get(Lead, lead_id)
    if lead is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lead")
    ensure_lead_access(db, user, lead)
    if lead.converted_at is not None:
        raise HTTPException(status_code=409, detail="Lead đã chuyển đổi và đã được khóa")
    note = str(payload.get("content", payload.get("note", ""))).strip()
    if not note or len(note) > 2000:
        raise HTTPException(status_code=422, detail="Ghi chú lịch sử phải có từ 1 đến 2000 ký tự")
    interaction_type = str(payload.get("interaction_type", "message"))
    if interaction_type not in {"call", "message", "meeting"}:
        raise HTTPException(status_code=422, detail="Loại tương tác phải là call, message hoặc meeting")
    try:
        follow_up_at = None
        if payload.get("follow_up_at"):
            dt = datetime.fromisoformat(str(payload["follow_up_at"]).replace("Z", "+00:00"))
            follow_up_at = dt.astimezone(timezone.utc).replace(tzinfo=None) if dt.tzinfo else dt
    except ValueError as error:
        raise HTTPException(status_code=422, detail="Ngày nhắc lại không hợp lệ") from error

    created_at = None
    if payload.get("created_at"):
        try:
            dt = datetime.fromisoformat(str(payload["created_at"]).replace("Z", "+00:00"))
            created_at = dt.astimezone(timezone.utc).replace(tzinfo=None) if dt.tzinfo else dt
        except ValueError:
            created_at = None

    record_lead_interaction(
        db, lead, user, interaction_type, note,
        interaction_type=interaction_type, follow_up_at=follow_up_at,
        created_at=created_at,
    )
    db.commit()
    return {"message": "Đã lưu ghi chú trao đổi", "interaction_type": interaction_type, "follow_up_at": follow_up_at.isoformat() if follow_up_at else None}


@app.post("/auth/login", response_model=LoginResponse)
def login(request: LoginRequest, db: DbSession) -> LoginResponse:
    email = request.email.strip().lower()
    user = db.scalar(select(User).where(User.email == email))
    now = datetime.now(timezone.utc)
    login_attempt = db.get(LoginAttempt, email)
    locked_until = login_attempt.locked_until if login_attempt else None
    if locked_until is not None and locked_until.tzinfo is None:
        locked_until = locked_until.replace(tzinfo=timezone.utc)
    if user and user.locked_until:
        account_locked_until = user.locked_until
        if account_locked_until.tzinfo is None:
            account_locked_until = account_locked_until.replace(tzinfo=timezone.utc)
        if locked_until is None or account_locked_until > locked_until:
            locked_until = account_locked_until
    if locked_until is not None and locked_until > now:
        remaining_seconds = max(1, int((locked_until - now).total_seconds()) + 1)
        remaining = max(1, (remaining_seconds + 59) // 60)
        raise HTTPException(
            status_code=status.HTTP_423_LOCKED,
            detail=f"Tài khoản tạm khóa. Vui lòng thử lại sau {remaining} phút",
            headers={"Retry-After": str(remaining_seconds)},
        )
    if locked_until is not None:
        if login_attempt:
            login_attempt.failed_attempts = 0
            login_attempt.locked_until = None
        if user:
            user.locked_until = None
            user.failed_login_attempts = 0
    valid_password = (
        verify_password(request.password, user.password_hash)
        if user
        else verify_password(request.password, "invalid$invalid")
    )
    if not valid_password:
        if login_attempt is None:
            login_attempt = LoginAttempt(
                email=email,
                failed_attempts=user.failed_login_attempts if user else 0,
            )
            db.add(login_attempt)
        login_attempt.failed_attempts += 1
        if user:
            user.failed_login_attempts = login_attempt.failed_attempts
        if login_attempt.failed_attempts >= 5:
            locked_until = now + timedelta(minutes=15)
            login_attempt.locked_until = locked_until.replace(tzinfo=None)
            if user:
                user.locked_until = locked_until.replace(tzinfo=None)
                user.lock_reason = "Tự động khóa sau 5 lần đăng nhập sai liên tiếp"
                user.auth_version += 1
        db.commit()
        if login_attempt.failed_attempts >= 5:
            raise HTTPException(
                status_code=status.HTTP_423_LOCKED,
                detail="Tài khoản tạm khóa 15 phút sau 5 lần đăng nhập sai",
                headers={"Retry-After": str(15 * 60)},
            )
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Email hoặc mật khẩu không đúng",
        )
    if user is None:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Email hoặc mật khẩu không đúng",
        )
    if not user.is_active:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Tài khoản của bạn đã bị khóa, xin liên hệ với quản trị viên.",
        )
    if login_attempt:
        login_attempt.failed_attempts = 0
        login_attempt.locked_until = None
    user.failed_login_attempts = 0
    user.locked_until = None
    db.commit()
    return LoginResponse(access_token=encode_token(user), user=user_response(db, user))


@app.post("/auth/change-password")
def change_password(
    request: PasswordChangeRequest, user: CurrentUser, db: DbSession
) -> dict[str, str | UserResponse]:
    if not verify_password(request.current_password, user.password_hash):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Mật khẩu hiện tại không đúng",
        )
    user.password_hash = hash_password(request.new_password)
    user.must_change_password = False
    user.auth_version += 1
    db.commit()
    return {
        "message": "Đổi mật khẩu thành công",
        "access_token": encode_token(user),
        "user": user_response(db, user),
    }


@app.post("/auth/forgot-password")
def forgot_password(request: ForgotPasswordRequest, db: DbSession) -> dict[str, str]:
    generic_response = {
        "message": (
            "Nếu email tồn tại, liên kết sẽ được gửi có hiệu lực trong 30 phút và chỉ dùng một lần."
        )
    }
    user = db.scalar(select(User).where(User.email == request.email))
    if user is None or not user.is_active:
        return generic_response
    token = secrets.token_urlsafe(32)
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    token_record = PasswordResetToken(
        token_hash=hashlib.sha256(token.encode()).hexdigest(),
        user_id=user.id,
        expires_at=now + timedelta(minutes=30),
    )
    db.execute(
        delete(PasswordResetToken).where(
            PasswordResetToken.user_id == user.id,
            (PasswordResetToken.used_at.is_not(None))
            | (PasswordResetToken.expires_at <= now),
        )
    )
    db.add(token_record)
    db.commit()
    message = reset_email(user.email, token)
    try:
        delivered = send_email(message)
    except (OSError, smtplib.SMTPException):
        token_record.used_at = datetime.now(timezone.utc).replace(tzinfo=None)
        db.commit()
        logger.exception("Unable to send password reset email")
        return generic_response
    if not delivered:
        logger.info("Development password reset link: %s", message.get_content().splitlines()[1])
    return generic_response


@app.get("/auth/reset-password/validate")
def validate_reset_password_token(
    token: Annotated[str, Query(min_length=1)],
    db: DbSession,
) -> dict[str, str]:
    token_hash = hashlib.sha256(token.encode()).hexdigest()
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    token_record = db.scalar(
        select(PasswordResetToken).where(
            PasswordResetToken.token_hash == token_hash,
            PasswordResetToken.used_at.is_(None),
            PasswordResetToken.expires_at > now,
        )
    )
    user = db.get(User, token_record.user_id) if token_record else None
    if token_record is None or user is None or not user.is_active:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Liên kết đặt lại mật khẩu không còn hiệu lực",
        )
    return {"email": user.email}


@app.post("/auth/reset-password")
def reset_password(
    request: ResetPasswordRequest, db: DbSession
) -> dict[str, str]:
    token_hash = hashlib.sha256(request.token.encode()).hexdigest()
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    token_record = db.scalar(
        select(PasswordResetToken).where(
            PasswordResetToken.token_hash == token_hash,
            PasswordResetToken.used_at.is_(None),
            PasswordResetToken.expires_at > now,
        )
    )
    user = db.get(User, token_record.user_id) if token_record else None
    if token_record is None or user is None or not user.is_active:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Liên kết đặt lại mật khẩu không còn hiệu lực",
        )
    claimed = db.execute(
        update(PasswordResetToken)
        .where(
            PasswordResetToken.id == token_record.id,
            PasswordResetToken.used_at.is_(None),
            PasswordResetToken.expires_at > now,
        )
        .values(used_at=now)
    )
    if claimed.rowcount != 1:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Liên kết đặt lại mật khẩu không còn hiệu lực",
        )
    user.password_hash = hash_password(request.new_password)
    user.must_change_password = False
    user.auth_version += 1
    db.commit()
    return {"message": "Đặt lại mật khẩu thành công"}


@app.post("/auth/logout")
def logout(user: CurrentUser, db: DbSession) -> dict[str, str]:
    user.auth_version += 1
    db.commit()
    return {"message": "Đã đăng xuất"}


@app.post("/auth/refresh", response_model=LoginResponse)
def refresh_session(
    credentials: CurrentCredentials, user: CurrentUser, db: DbSession
) -> LoginResponse:
    if credentials is None:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Phiên đăng nhập đã hết hạn",
        )
    payload = decode_token(credentials.credentials)
    return LoginResponse(
        access_token=encode_token(user, payload["started"]),
        user=user_response(db, user),
    )


@app.get("/dashboard/summary")
def dashboard_summary(user: CurrentUser, db: DbSession) -> dict[str, int | str]:
    if not permission_allows(assigned_roles(db, user), "users", "read"):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Bạn không có quyền xem bảng tổng quan quản trị",
        )
    users = list(db.scalars(select(User)).all())
    return {
        "greeting": user.full_name,
        "total_users": len(users),
        "active_users": sum(item.is_active for item in users),
        "locked_users": sum(not item.is_active for item in users),
    }


@app.get("/users", response_model=UserPage)
def list_users(
    user: Annotated[User, Depends(require_permission("users", "read"))],
    db: DbSession,
    q: str | None = None,
    role: RoleCode | None = None,
    active: bool | None = None,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
) -> UserPage:
    statement = select(User)
    if q:
        search = f"%{q.strip()}%"
        statement = statement.where(
            (User.full_name.ilike(search))
            | (User.email.ilike(search))
            | (User.phone.ilike(search))
        )
    if role:
        matching_ids = select(UserRoleAssignment.user_id).where(
            UserRoleAssignment.role == role
        )
        statement = statement.where(User.id.in_(matching_ids))
    if active is not None:
        statement = statement.where(User.is_active == active)
    total = db.scalar(select(text("count(*)")).select_from(statement.subquery())) or 0
    users = list(
        db.scalars(
            statement.order_by(User.created_at.desc(), User.id.desc())
            .offset((page - 1) * page_size)
            .limit(page_size)
        ).all()
    )
    return UserPage(
        items=[user_response(db, item) for item in users],
        total=total,
        page=page,
        page_size=page_size,
    )


@app.post("/users", response_model=UserResponse, status_code=status.HTTP_201_CREATED)
def create_user(
    request: UserCreate,
    _: Annotated[User, Depends(require_permission("users", "full"))],
    db: DbSession,
) -> UserResponse:
    if db.scalar(select(User).where(User.email == request.email)) is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT, detail="Email đã tồn tại"
        )
    temporary_password = f"{secrets.token_urlsafe(12)}A1!"
    user = User(
        full_name=request.full_name.strip(),
        email=request.email,
        phone=request.phone.strip() if request.phone else None,
        password_hash=hash_password(temporary_password),
        role=request.roles[0],
        must_change_password=True,
    )
    db.add(user)
    db.flush()
    replace_roles(db, user, request.roles)
    replace_class_assignments(db, user, request.assigned_classes)
    try:
        delivered = send_email(
            activation_email(user.email, user.full_name, temporary_password)
        )
    except (OSError, smtplib.SMTPException) as error:
        db.rollback()
        logger.exception("Unable to send account activation email")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không gửi được email kích hoạt; tài khoản chưa được tạo",
        ) from error
    if not delivered:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Chưa cấu hình email kích hoạt; tài khoản chưa được tạo",
        )
    db.commit()
    db.refresh(user)
    return user_response(db, user)


# Register Sprint 3/4 models and routes before FastAPI runs its startup migration.
import sprint34  # noqa: E402
import sprint56  # noqa: E402
import sprint78  # noqa: E402





@app.patch("/users/{user_id}", response_model=UserResponse)
def update_user(
    user_id: int,
    request: UserUpdate,
    _: Annotated[User, Depends(require_permission("users", "full"))],
    db: DbSession,
) -> UserResponse:
    user = db.get(User, user_id)
    if user is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy tài khoản"
        )
    duplicate = db.scalar(
        select(User).where(User.email == request.email, User.id != user_id)
    )
    if duplicate is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT, detail="Email đã tồn tại"
        )
    user.full_name = request.full_name.strip()
    user.email = request.email
    user.phone = request.phone.strip() if request.phone else None
    replace_class_assignments(db, user, request.assigned_classes)
    db.commit()
    db.refresh(user)
    return user_response(db, user)


@app.patch("/users/{user_id}/status", response_model=UserResponse)
def update_user_status(
    user_id: int,
    request: UserStatusUpdate,
    actor: Annotated[User, Depends(require_permission("users", "full"))],
    db: DbSession,
) -> UserResponse:
    user = db.get(User, user_id)
    if user is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy tài khoản"
        )
    if actor.id == user.id and not request.active:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Bạn không thể tự khóa tài khoản đang sử dụng",
        )
    reason = (request.reason or "").strip()
    if not request.active and not reason:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="Vui lòng nhập lý do khóa tài khoản",
        )
    user.is_active = request.active
    user.lock_reason = None if request.active else reason
    if request.active:
        login_attempt = db.get(LoginAttempt, user.email.lower())
        if login_attempt:
            login_attempt.failed_attempts = 0
            login_attempt.locked_until = None
        user.failed_login_attempts = 0
        user.locked_until = None
    if not request.active:
        user.auth_version += 1
    db.commit()
    db.refresh(user)
    return user_response(db, user)


@app.delete("/users/{user_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_user(
    user_id: int,
    actor: Annotated[User, Depends(require_permission("users", "full"))],
    db: DbSession,
) -> Response:
    user = db.get(User, user_id)
    if user is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy tài khoản"
        )
    if actor.id == user.id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Bạn không thể xóa tài khoản đang sử dụng",
        )
    target_roles = assigned_roles(db, user)
    if user.is_active and "admin" in target_roles:
        active_admin_count = db.scalar(
            select(func.count(User.id))
            .join(UserRoleAssignment, UserRoleAssignment.user_id == User.id)
            .where(UserRoleAssignment.role == "admin", User.is_active.is_(True))
        )
        if active_admin_count == 1:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Không thể xóa quản trị viên hoạt động cuối cùng",
            )
    db.execute(delete(UserRoleAssignment).where(UserRoleAssignment.user_id == user_id))
    db.execute(delete(UserClassAssignment).where(UserClassAssignment.user_id == user_id))
    db.execute(delete(PasswordResetToken).where(PasswordResetToken.user_id == user_id))
    db.execute(
        update(Lead)
        .where(Lead.assigned_to_user_id == user_id)
        .values(assigned_to_user_id=None)
    )
    db.delete(user)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@app.patch("/users/{user_id}/role", response_model=UserResponse)
def update_user_role(
    user_id: int,
    request: UserRoleUpdate,
    actor: Annotated[User, Depends(require_permission("users", "full"))],
    db: DbSession,
) -> UserResponse:
    user = db.get(User, user_id)
    if user is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy tài khoản"
        )
    if actor.id == user.id and "admin" in assigned_roles(db, user) and "admin" not in request.roles:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Bạn không thể tự thu hồi vai trò quản trị của mình",
        )
    replace_roles(db, user, request.roles)
    db.commit()
    db.refresh(user)
    return user_response(db, user)
